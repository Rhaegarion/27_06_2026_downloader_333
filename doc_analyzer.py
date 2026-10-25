"""
doc_analyzer.py - Nexus AI document analyzer (backend only)

Upload (PDF / DOCX / DOC / CSV / XLSX / images) -> short English summary report
(key persons, key events, important details) -> follow-up Q&A for the chatbot.

Pipeline
  1. Extract  : native text first; scanned pages and legacy-font pages (Bijoy/SutonnyMJ,
                Kruti Dev ...) are rendered and sent to Tesseract.
  2. Route    : low-confidence pages, Urdu/Pashto (Nastaliq) pages and pictures -> Gemma vision.
  3. Map      : Gemma reads the ORIGINAL language and returns English JSON facts per chunk
                (no full translation pass -> far fewer output tokens -> much faster).
  4. Reduce   : entities merged in Python, one final Gemma call writes the report.
  5. Index    : (background, only for large documents) EmbeddingGemma + Qdrant for Q&A.

Public API (everything the frontend needs)
  start_analysis(file_bytes, filename, user_prompt=None) -> doc_id      # non-blocking
  get_status(doc_id) -> dict                                            # poll for progress
  get_result(doc_id) -> dict | None                                     # summary when done
  analyze_document(file_bytes, filename, user_prompt=None, progress_cb=None) -> dict  # blocking
  ask_document(doc_id, question, history=None) -> str                   # follow-up Q&A
  delete_document(doc_id) / cleanup_old_documents(max_age_days=7)
"""
from __future__ import annotations

import base64
import hashlib
import io
import json
import operator
import re
import shutil
import subprocess
import tempfile
import threading
import time
import traceback
import unicodedata
import uuid
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import pymupdf as fitz  # PyMuPDF
import pandas as pd
import pytesseract
import requests
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageSequence, ImageStat
from qdrant_client import QdrantClient
from qdrant_client.models import (Distance, FieldCondition, Filter, FilterSelector,
                                  MatchValue, PointStruct, VectorParams)

# =====================================================================================
# CONFIG - edit here
# =====================================================================================
OLLAMA_URL = "http://localhost:11434"
LLM_MODEL = "gemma4:31b"
EMBED_MODEL = "embeddinggemma:latest"
NUM_CTX = 32768          # KEEP EQUAL to the num_ctx the chatbot uses - a different value makes
                         # Ollama reload gemma4 every time the two features alternate.
OLLAMA_THINK = None      # set False if gemma4 "thinks" by default in your Ollama build (faster)
KEEP_ALIVE = "30m"
LLM_TIMEOUT = 900
PARALLEL_LLM = 4         # also set OLLAMA_NUM_PARALLEL=4 for the Ollama server
PARALLEL_OCR = 8         # Tesseract processes (CPU)

TESSERACT_CMD = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
SOFFICE_CMD = r"C:\Program Files\LibreOffice\program\soffice.exe"   # only for .doc / legacy fonts

WORK_DIR = Path(r"D:\GPT\Nexus1.0\DOCS\work")
# Separate local Qdrant store: local mode locks its folder, so it must NOT be qd_storage_call.
DOCS_QDRANT_PATH = r"D:\GPT\Nexus1.0\DOCS\qd_storage_docs"
DOCS_COLLECTION = "uploaded_docs"

OCR_DPI = 300
MIN_NATIVE_CHARS = 50        # fewer chars of real text on a PDF page -> treat as scanned
OCR_CONF_THRESHOLD = 75      # avg Tesseract word confidence below this -> Gemma vision
NASTALIQ_CHAR_MIN = 3        # Urdu/Pashto-only letters in OCR output -> Gemma vision
BLANK_STDDEV = 4             # near-uniform scanned page -> skipped
DEFAULT_OCR_LANGS = "eng+ben+hin"   # used only if Tesseract's osd.traineddata is missing
SCRIPT_TO_LANGS = {
    "Bengali": "ben+eng",
    "Devanagari": "hin+nep+eng",
    "Arabic": "ara+urd+pus",
    "Han": "chi_sim+chi_tra", "HanS": "chi_sim", "HanT": "chi_tra",
    "Latin": "eng",
}
LEGACY_FONT_HINTS = ("sutonny", "bijoy", "krutidev", "chanakya", "devlys", "shusha", "walkman")

MIN_IMAGE_SIDE = 150         # smaller pictures (icons, bullets) are ignored
DUP_HASH_DISTANCE = 4        # near-identical pictures (logos repeated per page) analysed once
MAX_IMAGES_PER_DOC = 60
VISION_MAX_SIDE = 1600

DOCX_PART_CHARS = 3000       # DOCX has no pages -> split into "parts" for citations
CHUNK_CHARS = 12000          # text per Gemma extraction call
SECTION_CHARS = 15000
FINAL_SUMMARIES_MAX_CHARS = 30000
MAX_PERSONS, MAX_EVENTS, MAX_FACTS, MAX_IMAGES_IN_REPORT = 60, 80, 80, 40
MAP_MAX_TOKENS, FINAL_MAX_TOKENS, QA_MAX_TOKENS = 3000, 4000, 1500

FULL_CONTEXT_CHARS = 60000   # docs smaller than this: Q&A uses the whole text, no embeddings
EMBED_CHUNK_CHARS = 1200
EMBED_BATCH = 32
TOP_K = 8
KEYWORD_K = 4
TABLE_SAMPLE_ROWS = 20
PROMPT_VERSION = "v1"        # bump to invalidate cached per-page results after prompt edits

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
TABLE_EXTS = {".csv", ".xlsx", ".xlsm", ".xls"}
URDU_PASHTO_CHARS = set("ٹڈڑںےۓھټډړږښځڅۍېګڼ")
_CITE_PREFIX = {"page": "p.", "part": "part", "sheet": "sheet"}

if TESSERACT_CMD and Path(TESSERACT_CMD).exists():
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD

# =====================================================================================
# PROMPTS
# =====================================================================================
SYSTEM_PROMPT = ("You are a careful document analyst. You extract facts exactly as stated "
                 "in the source, never invent information, and write all output in English.")

EXTRACT_RULES = """Rules:
- Read the source in whatever language it is written (Bengali, Hindi, Nepali, Chinese, Urdu, Pashto, Arabic, English, ...). Write ALL output in English.
- Keep names, dates, numbers, phone numbers, IDs, vehicle numbers and amounts exactly as written. For every person give "name" as written in the source and "name_en" in English letters.
- Use only information present in the source. Do not guess. Use "" or [] when something is not stated.
- "page" must be the <<label>> number where the item appears.
- Return ONLY one JSON object, no markdown."""

EXTRACT_SCHEMA = """{
  "persons": [{"name": "", "name_en": "", "role": "", "details": "what the source says about them, 1 sentence", "page": 0}],
  "events": [{"date": "", "description": "", "location": "", "persons": [""], "page": 0}],
  "organisations": [""],
  "locations": [""],
  "key_facts": [{"fact": "numbers, IDs, phone numbers, vehicles, amounts, decisions, findings", "page": 0}],
  "page_summaries": [{"page": 0, "summary": "2-3 sentences in English"}],
  "languages": ["language(s) of the source text"]
}"""

TEXT_PROMPT = """Below is text from <<label>>(s) <<pages>> of an uploaded document. Lines like [<<LABEL>> 3] mark where each <<label>> starts. Text from scanned pages may contain OCR errors; correct obvious ones only when certain.

<<rules>>

JSON format:
<<schema>>

DOCUMENT TEXT:
<<text>>"""

VISION_PAGE_PROMPT = """This image is <<label>> <<page>> of a scanned or photographed document. It may be written in Urdu/Pashto (Nastaliq), Arabic, Bengali, Hindi, Nepali, Chinese or another language. Read it carefully.

<<rules>>
Also add these keys to the JSON:
  "key_text_en": English translation of the most important text on the page (at most 200 words),
  "visual_notes": photos, maps, stamps, signatures or handwriting on the page (1-2 sentences).
Do not identify people from their faces; name a person only if text on the page names them.

JSON format:
<<schema>>
<<ocr_hint>>"""

IMAGE_PROMPT = """This picture appears on <<label>> <<page>> of an uploaded document. Describe it factually for an analyst.
Rules:
- English only. Do not guess.
- Transcribe and translate any visible text (captions, signs, documents, number plates, stamps).
- Do not identify people from their faces; name a person only if a caption or visible text names them. You may describe people (how many, actions, clothing, uniforms).
- Return ONLY this JSON:
{"description": "2-4 sentences", "visible_text_en": "", "persons": [{"name": "", "name_en": "", "role": "", "details": "", "page": <<page>>}], "locations": [""], "key_facts": [{"fact": "", "page": <<page>>}]}"""

COMPRESS_PROMPT = """Condense these <<label>> summaries from one document into a section summary of 6-10 sentences in English. Keep every important name, date, number and place, and keep the references in brackets like [p. 12].

<<text>>"""

FINAL_PROMPT = """You are writing a short summary report of an uploaded document for a busy officer.
Below is material extracted from the document: summaries per <<label>>, people, events, facts and pictures, with <<label>> numbers.

Write the report as ONE JSON object:
{
  "title": "short descriptive title of the document",
  "document_type": "e.g. assessment report, letter, situation report, spreadsheet",
  "overview": "4-6 sentences: what the document is, who it concerns, the main findings",
  "key_persons": [{"name": "English name (original script in brackets if available)", "role": "", "involvement": "1-2 sentences", "pages": [1]}],
  "key_events": [{"date": "", "event": "", "persons": [""], "location": "", "pages": [1]}],
  "important_details": [{"detail": "", "pages": [1]}],
  "answer_to_user_request": ""
}
Rules:
- Use only the material. Do not invent anything.
- Merge entries that clearly refer to the same person (spelling or transliteration variants).
- key_persons: at most 15, most important first. key_events: at most 15, chronological where dates allow. important_details: at most 12 (numbers, IDs, phone numbers, vehicles, amounts, decisions, notable pictures or tables).
- answer_to_user_request: answer the user's request below using the material, or "" if there is no request.
- English only. Return ONLY the JSON object.

USER REQUEST: <<user_prompt>>

MATERIAL:
<<material>>"""

QA_PROMPT = """You answer questions about ONE uploaded document ("<<filename>>") using ONLY the material below.
- Cite sources like (<<cite_example>>).
- If the answer is not in the material, say clearly that the document does not say.
- Answer in English unless the user writes in another language. Give names in English; add the original script in brackets the first time if available.
- Be concise and factual.

REPORT SUMMARY:
<<summary>>
<<computed>>
DOCUMENT MATERIAL:
<<context>>

CONVERSATION SO FAR:
<<history>>

QUESTION: <<question>>"""

TABLE_PLAN_PROMPT = """You plan a computation on a spreadsheet to answer a question.
Sheets and columns:
<<schema>>

Question: <<question>>

If answering needs filtering, counting, summing, averaging or ranking rows, return a plan; otherwise set "needs_computation" to false.
Return ONLY JSON:
{"needs_computation": true, "sheet": "", "filters": [{"column": "", "op": "==|!=|>|<|>=|<=|contains", "value": ""}], "group_by": [""], "agg": {"column": "", "func": "count|sum|mean|min|max|nunique"}, "select": [""], "sort_by": "", "ascending": false, "limit": 20}
Use exact column names from the list."""


# =====================================================================================
# SMALL HELPERS
# =====================================================================================
def _fill(template: str, **kw) -> str:
    for k, v in kw.items():
        template = template.replace(f"<<{k}>>", str(v))
    return template


def _read_json(path: Path, default=None):
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except Exception:
        return default


def _write_json(path: Path, obj) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(obj, ensure_ascii=False, indent=1, default=list), encoding="utf-8")
    tmp.replace(path)


def _log(d: Path, msg: str) -> None:
    try:
        with open(Path(d) / "log.txt", "a", encoding="utf-8") as f:
            f.write(f"{time.strftime('%Y-%m-%d %H:%M:%S')} {msg}\n")
    except Exception:
        pass


def _as_list(x) -> list:
    return x if isinstance(x, list) else []


def _norm_name(s: str) -> str:
    s = unicodedata.normalize("NFKC", str(s or "")).casefold()
    s = re.sub(r"[^\w\s]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _cite(label: str, pages) -> str:
    nums = set()
    for x in pages or []:
        try:
            nums.add(int(x))
        except (TypeError, ValueError):
            pass
    if not nums:
        return ""
    return f"{_CITE_PREFIX.get(label, 'p.')} {', '.join(str(n) for n in sorted(nums))}"


def _range_str(pages) -> str:
    pages = sorted(set(pages))
    if not pages:
        return ""
    if pages == list(range(pages[0], pages[-1] + 1)) and len(pages) > 1:
        return f"{pages[0]}-{pages[-1]}"
    return ", ".join(map(str, pages))


def _split(text: str, size: int, overlap: int = 0) -> list[str]:
    text = (text or "").strip()
    if len(text) <= size:
        return [text] if text else []
    out, i = [], 0
    while i < len(text):
        end = min(len(text), i + size)
        if end < len(text):
            lo = i + size // 2
            cut = max(text.rfind("\n", lo, end), text.rfind(". ", lo, end), text.rfind("। ", lo, end))
            if cut > i:
                end = cut + 1
        out.append(text[i:end].strip())
        if end >= len(text):
            break
        i = max(end - overlap, i + 1)
    return [o for o in out if o]


def _parse_json(raw: str):
    if not raw:
        return None
    s = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw.strip())
    try:
        return json.loads(s)
    except Exception:
        pass
    a, b = s.find("{"), s.rfind("}")
    if a != -1 and b > a:
        try:
            return json.loads(s[a:b + 1])
        except Exception:
            return None
    return None


def _table_to_md(rows) -> str:
    clean = []
    for r in rows or []:
        cells = [str(c or "").strip().replace("\n", " ").replace("|", "/") for c in r]
        if any(cells):
            clean.append(cells)
    if not clean:
        return ""
    width = max(len(r) for r in clean)
    clean = [r + [""] * (width - len(r)) for r in clean]
    lines = ["| " + " | ".join(clean[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in clean[1:]]
    return "\n".join(lines)


def _check_id(doc_id: str) -> str:
    if not re.fullmatch(r"[0-9a-f]{16}", str(doc_id or "")):
        raise ValueError("Invalid doc_id")
    return doc_id


def _make_doc_id(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()[:16]


def _doc_dir(doc_id: str) -> Path:
    d = WORK_DIR / doc_id
    d.mkdir(parents=True, exist_ok=True)
    return d


# =====================================================================================
# STATUS (for progress bars)
# =====================================================================================
_STATUS: dict[str, dict] = {}
_STATUS_LOCK = threading.Lock()
_RUNNING: set[str] = set()
_INDEXING: set[str] = set()


def _set_status(doc_id: str, **kw) -> None:
    with _STATUS_LOCK:
        st = _STATUS.setdefault(doc_id, {"state": "queued", "stage": "", "done": 0, "total": 0,
                                         "message": "", "index": "pending"})
        st.update(kw)
        st["updated"] = time.time()
        snap = dict(st)
    try:
        _write_json(_doc_dir(doc_id) / "status.json", snap)
    except Exception:
        pass


def get_status(doc_id: str) -> dict:
    """{state: queued|running|done|error, stage, done, total, message, index: pending|building|ready|error}"""
    _check_id(doc_id)
    with _STATUS_LOCK:
        if doc_id in _STATUS:
            return dict(_STATUS[doc_id])
    return _read_json(WORK_DIR / doc_id / "status.json", {"state": "unknown"})


def get_result(doc_id: str) -> dict | None:
    return _read_json(WORK_DIR / _check_id(doc_id) / "result.json")


# =====================================================================================
# OLLAMA (Gemma 4 + EmbeddingGemma)
# =====================================================================================
_LLM_SEM = threading.BoundedSemaphore(PARALLEL_LLM)


def _chat(prompt: str, images: list[bytes] | None = None, json_mode: bool = False,
          max_tokens: int = MAP_MAX_TOKENS) -> str:
    msg = {"role": "user", "content": prompt}
    if images:
        msg["images"] = [base64.b64encode(b).decode("ascii") for b in images]
    payload = {
        "model": LLM_MODEL,
        "messages": [{"role": "system", "content": SYSTEM_PROMPT}, msg],
        "stream": False,
        "keep_alive": KEEP_ALIVE,
        "options": {"temperature": 0.1, "num_ctx": NUM_CTX, "num_predict": max_tokens},
    }
    if json_mode:
        payload["format"] = "json"
    if OLLAMA_THINK is not None:
        payload["think"] = OLLAMA_THINK
    last_err = None
    for attempt in range(2):
        try:
            with _LLM_SEM:
                r = requests.post(f"{OLLAMA_URL}/api/chat", json=payload, timeout=LLM_TIMEOUT)
            r.raise_for_status()
            return ((r.json().get("message") or {}).get("content") or "").strip()
        except requests.RequestException as e:
            last_err = e
            time.sleep(3)
    raise RuntimeError(f"Ollama chat failed: {last_err}")


def _chat_json(prompt: str, images: list[bytes] | None = None,
               max_tokens: int = MAP_MAX_TOKENS) -> dict:
    for _ in range(2):
        data = _parse_json(_chat(prompt, images, json_mode=True, max_tokens=max_tokens))
        if isinstance(data, dict):
            return data
        prompt += "\n\nYour previous answer was not valid JSON. Return ONLY valid JSON."
    return {}


def _embed(texts: list[str], is_query: bool = False) -> list[list[float]]:
    # EmbeddingGemma expects different prefixes for queries and documents.
    pre = "task: search result | query: " if is_query else "title: none | text: "
    texts = [pre + t for t in texts]
    out: list[list[float]] = []
    for i in range(0, len(texts), EMBED_BATCH):
        r = requests.post(f"{OLLAMA_URL}/api/embed",
                          json={"model": EMBED_MODEL, "input": texts[i:i + EMBED_BATCH],
                                "keep_alive": KEEP_ALIVE},
                          timeout=LLM_TIMEOUT)
        r.raise_for_status()
        out.extend(r.json()["embeddings"])
    return out


# =====================================================================================
# QDRANT (separate local store for uploaded documents)
# =====================================================================================
_qdrant: QdrantClient | None = None
_qdrant_lock = threading.Lock()


def _get_qdrant() -> QdrantClient:
    global _qdrant
    if _qdrant is None:
        Path(DOCS_QDRANT_PATH).mkdir(parents=True, exist_ok=True)
        _qdrant = QdrantClient(path=DOCS_QDRANT_PATH)
    return _qdrant


def _doc_filter(doc_id: str) -> Filter:
    return Filter(must=[FieldCondition(key="doc_id", match=MatchValue(value=doc_id))])


# =====================================================================================
# IMAGES
# =====================================================================================
def _to_rgb(img: Image.Image) -> Image.Image:
    try:
        return img.convert("RGB")
    except Exception:
        return img.convert("L").convert("RGB")


def _to_jpeg(img: Image.Image, max_side: int = VISION_MAX_SIDE) -> bytes:
    img = _to_rgb(img)
    w, h = img.size
    s = max(w, h)
    if s > max_side:
        img = img.resize((max(1, int(w * max_side / s)), max(1, int(h * max_side / s))), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=90)
    return buf.getvalue()


def _vision_bytes(path: str) -> bytes:
    with Image.open(path) as img:
        return _to_jpeg(img)


def _dhash(img: Image.Image) -> int:
    g = img.convert("L").resize((9, 8), Image.LANCZOS)
    px = list(g.getdata())
    bits = 0
    for r in range(8):
        for c in range(8):
            bits = (bits << 1) | int(px[r * 9 + c] > px[r * 9 + c + 1])
    return bits


class _ImageKeeper:
    """Drops tiny, blank and duplicate pictures (logos/stamps repeated on every page)."""

    def __init__(self, d: Path):
        self.dir = d / "images"
        self.dir.mkdir(parents=True, exist_ok=True)
        self.hashes: list[int] = []
        self.count = 0
        self.skipped_over_limit = 0

    def add(self, raw: bytes) -> str | None:
        try:
            img = Image.open(io.BytesIO(raw))
            img.load()
        except Exception:
            return None
        if min(img.size) < MIN_IMAGE_SIDE:
            return None
        try:
            if ImageStat.Stat(img.convert("L")).stddev[0] < 8:
                return None
        except Exception:
            pass
        h = _dhash(img)
        if any(bin(h ^ k).count("1") <= DUP_HASH_DISTANCE for k in self.hashes):
            return None
        if self.count >= MAX_IMAGES_PER_DOC:
            self.skipped_over_limit += 1
            return None
        self.hashes.append(h)
        self.count += 1
        path = self.dir / f"img_{self.count:04d}.jpg"
        path.write_bytes(_to_jpeg(img))
        return str(path)


# =====================================================================================
# EXTRACTION
# =====================================================================================
def _new_page(n: int) -> dict:
    return {"page": n, "text": "", "tables": [], "images": [], "scan": None,
            "source": "native", "route": "text", "ocr_conf": None, "script": None}


def _is_legacy_font(name: str) -> bool:
    n = str(name or "").split("+")[-1].casefold()
    n = re.split(r"[-,]", n)[0].replace(" ", "")
    return n.endswith("mj") or any(h in n for h in LEGACY_FONT_HINTS)


def _has_legacy_font(page) -> bool:
    try:
        return any(_is_legacy_font(f[3]) for f in page.get_fonts())
    except Exception:
        return False


def _pdf_tables(page) -> list[str]:
    out = []
    try:
        for t in page.find_tables().tables:
            md = _table_to_md(t.extract())
            if md:
                out.append(md)
    except Exception:
        pass
    return out


def _extract_pdf(data: bytes, d: Path, keeper: _ImageKeeper) -> list[dict]:
    pages, seen_xrefs = [], set()
    scans = d / "scans"
    scans.mkdir(exist_ok=True)
    with fitz.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc, start=1):
            p = _new_page(i)
            text = (page.get_text("text") or "").strip()
            if len(text) >= MIN_NATIVE_CHARS and not _has_legacy_font(page):
                p["text"] = text
                p["tables"] = _pdf_tables(page)
                for info in page.get_images(full=True):
                    xref = info[0]
                    if xref in seen_xrefs:
                        continue
                    seen_xrefs.add(xref)
                    try:
                        raw = doc.extract_image(xref)["image"]
                    except Exception:
                        continue
                    path = keeper.add(raw)
                    if path:
                        p["images"].append(path)
            else:  # scanned page or garbled legacy-font text -> OCR
                path = scans / f"p{i:04d}.png"
                page.get_pixmap(dpi=OCR_DPI).save(str(path))
                p["scan"], p["source"] = str(path), "ocr"
            pages.append(p)
    return pages


def _iter_docx_blocks(doc):
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)


def _docx_has_legacy_fonts(doc) -> bool:
    for para in doc.paragraphs:
        names = [r.font.name for r in para.runs]
        try:
            names.append(para.style.font.name)
        except Exception:
            pass
        if any(n and _is_legacy_font(n) for n in names):
            return True
    return False


def _extract_docx(doc, keeper: _ImageKeeper) -> list[dict]:
    pages: list[dict] = []
    buf: list[str] = []
    imgs: list[str] = []
    size = 0

    def flush():
        nonlocal buf, imgs, size
        if buf or imgs:
            p = _new_page(len(pages) + 1)
            p["text"], p["images"] = "\n".join(buf).strip(), list(imgs)
            pages.append(p)
        buf, imgs, size = [], [], 0

    for block in _iter_docx_blocks(doc):
        if isinstance(block, Paragraph):
            for rid in block._p.xpath(".//a:blip/@r:embed"):
                part = doc.part.related_parts.get(rid)
                if part is not None:
                    path = keeper.add(part.blob)
                    if path:
                        imgs.append(path)
            t = block.text.strip()
            if t:
                buf.append(t)
                size += len(t)
        else:
            md = _table_to_md([[c.text for c in row.cells] for row in block.rows])
            if md:
                buf.append("[TABLE]\n" + md)
                size += len(md)
        if size >= DOCX_PART_CHARS:
            flush()
    flush()
    return pages


def _soffice_available() -> bool:
    return bool(SOFFICE_CMD) and Path(SOFFICE_CMD).exists()


def _soffice_convert(data: bytes, suffix: str, target: str) -> bytes:
    if not _soffice_available():
        raise RuntimeError(f"LibreOffice not found at {SOFFICE_CMD} (needed for {suffix} files)")
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / f"input{suffix}"
        src.write_bytes(data)
        subprocess.run([SOFFICE_CMD, "--headless", "--convert-to", target, "--outdir", td, str(src)],
                       check=True, timeout=300, capture_output=True)
        return (Path(td) / f"input.{target}").read_bytes()


def _read_csv(data: bytes) -> pd.DataFrame:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            return pd.read_csv(io.BytesIO(data), encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return pd.read_csv(io.BytesIO(data), encoding="latin-1")


def _profile_df(name: str, df: pd.DataFrame) -> str:
    lines = [f"### Sheet: {name}", f"Rows: {len(df)}, Columns: {len(df.columns)}", "Columns:"]
    for col in df.columns:
        s = df[col]
        line = f"- {col} ({s.dtype}); non-empty={int(s.notna().sum())}; unique={int(s.nunique(dropna=True))}"
        try:
            if pd.api.types.is_numeric_dtype(s) and s.notna().any():
                line += f"; min={s.min()}, max={s.max()}, mean={round(float(s.mean()), 2)}"
            else:
                top = s.dropna().astype(str).value_counts().head(5)
                if len(top):
                    line += "; top: " + ", ".join(f"{k} ({v})" for k, v in top.items())
        except Exception:
            pass
        lines.append(line)
    lines.append(f"Sample rows (first {TABLE_SAMPLE_ROWS}):")
    lines.append(df.head(TABLE_SAMPLE_ROWS).to_csv(index=False))
    return "\n".join(lines)


def _extract_table_file(data: bytes, ext: str, d: Path) -> list[dict]:
    raw = {"Sheet1": _read_csv(data)} if ext == ".csv" else pd.read_excel(io.BytesIO(data), sheet_name=None)
    dfs = {}
    for name, df in raw.items():
        df = df.dropna(how="all").dropna(axis=1, how="all")
        df.columns = [str(c).strip() for c in df.columns]
        dfs[str(name)] = df
    pd.to_pickle(dfs, d / "tables.pkl")
    pages = []
    for i, (name, df) in enumerate(dfs.items(), start=1):
        p = _new_page(i)
        p["text"], p["sheet"] = _profile_df(name, df), name
        pages.append(p)
    return pages


def _extract_image_file(data: bytes, d: Path) -> list[dict]:
    scans = d / "scans"
    scans.mkdir(exist_ok=True)
    pages = []
    with Image.open(io.BytesIO(data)) as img:
        for i, frame in enumerate(ImageSequence.Iterator(img), start=1):
            path = scans / f"p{i:04d}.png"
            _to_rgb(frame).save(path)
            p = _new_page(i)
            p["scan"], p["source"] = str(path), "ocr"
            pages.append(p)
    return pages


def _extract(data: bytes, filename: str, d: Path):
    """Returns (pages, unit_label, kind, keeper)."""
    ext = Path(filename).suffix.lower()
    keeper = _ImageKeeper(d)
    if ext == ".pdf":
        return _extract_pdf(data, d, keeper), "page", "document", keeper
    if ext in (".doc", ".rtf", ".odt"):
        data, ext = _soffice_convert(data, ext, "docx"), ".docx"
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        if _docx_has_legacy_fonts(doc) and _soffice_available():
            pdf = _soffice_convert(data, ".docx", "pdf")   # legacy fonts -> render + OCR
            return _extract_pdf(pdf, d, keeper), "page", "document", keeper
        return _extract_docx(doc, keeper), "part", "document", keeper
    if ext in TABLE_EXTS:
        return _extract_table_file(data, ext, d), "sheet", "table", keeper
    if ext in IMAGE_EXTS:
        return _extract_image_file(data, d), "page", "document", keeper
    raise ValueError(f"Unsupported file type: {ext}")


# =====================================================================================
# OCR (Tesseract) + routing to Gemma vision
# =====================================================================================
_TESS_LANGS: set[str] | None = None


def _installed_langs() -> set[str]:
    global _TESS_LANGS
    if _TESS_LANGS is None:
        try:
            _TESS_LANGS = set(pytesseract.get_languages(config=""))
        except Exception:
            _TESS_LANGS = {"eng"}
    return _TESS_LANGS


def _langs(spec: str) -> str | None:
    have = [l for l in spec.split("+") if l in _installed_langs()]
    return "+".join(have) if have else None


def _ocr_page(p: dict) -> dict:
    img = Image.open(p["scan"])
    img.load()
    if ImageStat.Stat(img.convert("L")).stddev[0] < BLANK_STDDEV:
        p["route"] = "skip"
        return p

    if "osd" in _installed_langs():
        try:
            osd = pytesseract.image_to_osd(img, output_type=pytesseract.Output.DICT)
        except Exception:      # too little text for OSD -> photo / handwriting / odd page
            p["route"] = "vision"
            return p
        p["script"] = osd.get("script", "Unknown")
        rot = int(osd.get("rotate", 0) or 0)
        if rot and float(osd.get("orientation_conf", 0) or 0) > 2:
            img = img.rotate(-rot, expand=True)
            img.save(p["scan"])
        lang = _langs(SCRIPT_TO_LANGS.get(p["script"], "eng"))
    else:
        lang = _langs(DEFAULT_OCR_LANGS)
    if not lang:
        p["route"] = "vision"
        return p

    data = pytesseract.image_to_data(img, lang=lang, config="--psm 3",
                                     output_type=pytesseract.Output.DICT)
    lines: dict[tuple, list[str]] = {}
    weighted, total = 0.0, 0
    for i, w in enumerate(data["text"]):
        w = (w or "").strip()
        try:
            c = float(data["conf"][i])
        except (TypeError, ValueError):
            c = -1
        if not w or c < 0:
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        lines.setdefault(key, []).append(w)
        weighted += c * len(w)
        total += len(w)
    text = "\n".join(" ".join(ws) for ws in lines.values())
    conf = weighted / total if total else 0.0
    nastaliq = sum(ch in URDU_PASHTO_CHARS for ch in text) >= NASTALIQ_CHAR_MIN

    p["text"], p["ocr_conf"] = text, round(conf, 1)
    low = conf < OCR_CONF_THRESHOLD or len(text.strip()) < MIN_NATIVE_CHARS
    p["route"] = "vision" if (low or nastaliq) else "text"
    return p


# =====================================================================================
# MAP: one Gemma call per unit (text chunk / vision page / picture), cached on disk
# =====================================================================================
def _page_text(p: dict) -> str:
    parts = [p.get("text") or ""] + ["[TABLE]\n" + t for t in p.get("tables") or []]
    return "\n\n".join(x for x in parts if x.strip())


def _build_units(pages: list[dict], label: str) -> list[dict]:
    units: list[dict] = []
    buf: list[str] = []
    buf_pages: set[int] = set()
    size = 0
    tag = label.upper()

    def flush():
        nonlocal buf, buf_pages, size
        if buf:
            units.append({"kind": "text", "pages": sorted(buf_pages), "text": "\n\n".join(buf)})
        buf, buf_pages, size = [], set(), 0

    for p in pages:
        route = p.get("route", "text")
        if route == "text":
            for piece in _split(_page_text(p), CHUNK_CHARS):
                block = f"[{tag} {p['page']}]\n{piece}"
                if buf and size + len(block) > CHUNK_CHARS:
                    flush()
                buf.append(block)
                buf_pages.add(p["page"])
                size += len(block)
        elif route == "vision":
            units.append({"kind": "vision", "pages": [p["page"]], "image": p["scan"],
                          "ocr_text": (p.get("text") or "")[:2000]})
        for img in p.get("images") or []:
            units.append({"kind": "image", "pages": [p["page"]], "image": img})
    flush()
    return units


def _run_unit(u: dict, label: str, d: Path) -> dict:
    sig = hashlib.sha1(f"{PROMPT_VERSION}|{u['kind']}|{u['pages']}|{label}".encode())
    if u.get("text"):
        sig.update(u["text"].encode("utf-8", "ignore"))
    if u.get("image"):
        sig.update(Path(u["image"]).read_bytes())
    cache = d / "cache" / f"{sig.hexdigest()}.json"
    hit = _read_json(cache)
    if hit is not None:
        return hit

    page = u["pages"][0] if u["pages"] else 0
    rules = _fill(EXTRACT_RULES, label=label)
    if u["kind"] == "text":
        prompt = _fill(TEXT_PROMPT, label=label, LABEL=label.upper(), pages=_range_str(u["pages"]),
                       rules=rules, schema=EXTRACT_SCHEMA, text=u["text"])
        data = _chat_json(prompt)
    elif u["kind"] == "vision":
        hint = ("\nRough machine OCR of this page (may be badly wrong; use only as a hint):\n"
                + u["ocr_text"]) if u.get("ocr_text") else ""
        prompt = _fill(VISION_PAGE_PROMPT, label=label, page=page, rules=rules,
                       schema=EXTRACT_SCHEMA, ocr_hint=hint)
        data = _chat_json(prompt, images=[_vision_bytes(u["image"])])
    else:
        data = _chat_json(_fill(IMAGE_PROMPT, label=label, page=page),
                          images=[_vision_bytes(u["image"])])

    out = {"kind": u["kind"], "pages": u["pages"], "data": data or {}}
    if data:
        _write_json(cache, out)
    return out


# =====================================================================================
# REDUCE: merge entities, final report
# =====================================================================================
def _merge(results: list[dict]) -> dict:
    persons: dict[str, dict] = {}
    events: dict[tuple, dict] = {}
    facts: dict[str, dict] = {}
    orgs, locs, langs = Counter(), Counter(), Counter()
    summaries: dict[int, str] = {}
    images: list[dict] = []

    for r in results:
        data = r.get("data") or {}
        default_page = r["pages"][0] if r.get("pages") else None

        def pg(x):
            try:
                return int(x.get("page") or default_page)
            except (TypeError, ValueError, AttributeError):
                return default_page

        for p in _as_list(data.get("persons")):
            if not isinstance(p, dict):
                continue
            name_en = str(p.get("name_en") or p.get("name") or "").strip()
            if not name_en:
                continue
            e = persons.setdefault(_norm_name(name_en), {"name": name_en, "original": set(), "roles": set(),
                                                          "details": [], "pages": set(), "mentions": 0})
            if p.get("name") and str(p["name"]).strip() != name_en:
                e["original"].add(str(p["name"]).strip())
            if p.get("role"):
                e["roles"].add(str(p["role"]).strip())
            det = str(p.get("details") or "").strip()
            if det and det not in e["details"] and len(e["details"]) < 5:
                e["details"].append(det)
            if pg(p) is not None:
                e["pages"].add(pg(p))
            e["mentions"] += 1

        for ev in _as_list(data.get("events")):
            if not isinstance(ev, dict) or not ev.get("description"):
                continue
            k = (_norm_name(ev.get("date")), _norm_name(ev["description"])[:80])
            e = events.setdefault(k, {"date": str(ev.get("date") or ""), "description": str(ev["description"]),
                                      "location": str(ev.get("location") or ""),
                                      "persons": [str(x) for x in _as_list(ev.get("persons")) if x][:6],
                                      "pages": set()})
            if pg(ev) is not None:
                e["pages"].add(pg(ev))

        for f in _as_list(data.get("key_facts")):
            text = f.get("fact") if isinstance(f, dict) else f
            if not text:
                continue
            e = facts.setdefault(_norm_name(text)[:100], {"fact": str(text), "pages": set()})
            if isinstance(f, dict) and pg(f) is not None:
                e["pages"].add(pg(f))
            elif default_page is not None:
                e["pages"].add(default_page)

        orgs.update(str(x).strip() for x in _as_list(data.get("organisations")) if x)
        locs.update(str(x).strip() for x in _as_list(data.get("locations")) if x)
        langs.update(str(x).strip() for x in _as_list(data.get("languages")) if x)

        for s in _as_list(data.get("page_summaries")):
            if isinstance(s, dict) and s.get("summary"):
                try:
                    n = int(s.get("page") or default_page)
                except (TypeError, ValueError):
                    n = default_page
                if n is not None:
                    summaries[n] = (summaries.get(n, "") + " " + str(s["summary"])).strip()

        if r["kind"] == "vision" and data.get("visual_notes") and default_page is not None:
            summaries[default_page] = (summaries.get(default_page, "") + " Visual: "
                                       + str(data["visual_notes"])).strip()
        if r["kind"] == "image" and data.get("description"):
            images.append({"page": default_page, "description": str(data["description"])[:300],
                           "text": str(data.get("visible_text_en") or "")[:200]})

    return {"persons": persons, "events": list(events.values()), "facts": list(facts.values()),
            "orgs": orgs, "locs": locs, "langs": langs, "summaries": summaries, "images": images}


def _compress_summaries(lines: list[str], label: str) -> str:
    batches, cur, size = [], [], 0
    for ln in lines:
        if cur and size + len(ln) > SECTION_CHARS:
            batches.append("\n".join(cur))
            cur, size = [], 0
        cur.append(ln)
        size += len(ln)
    if cur:
        batches.append("\n".join(cur))
    with ThreadPoolExecutor(max_workers=PARALLEL_LLM) as ex:
        outs = list(ex.map(lambda b: _chat(_fill(COMPRESS_PROMPT, label=label, text=b),
                                           max_tokens=800), batches))
    return "\n\n".join(f"Section {i + 1}: {o}" for i, o in enumerate(outs))


def _fallback_report(merged: dict, persons: list, events: list, facts: list, filename: str) -> dict:
    summ = [s for _, s in sorted(merged["summaries"].items())]
    return {
        "title": filename, "document_type": "", "overview": " ".join(summ[:5]),
        "key_persons": [{"name": e["name"], "role": ", ".join(e["roles"]),
                         "involvement": (e["details"] or [""])[0], "pages": e["pages"]} for e in persons[:15]],
        "key_events": [{"date": e["date"], "event": e["description"], "persons": e["persons"],
                        "location": e["location"], "pages": e["pages"]} for e in events[:15]],
        "important_details": [{"detail": f["fact"], "pages": f["pages"]} for f in facts[:12]],
        "answer_to_user_request": "",
    }


def _make_report(merged: dict, label: str, n_units: int, filename: str, user_prompt: str) -> dict:
    lines = [f"[{_cite(label, [p])}] {s}" for p, s in sorted(merged["summaries"].items())]
    summaries_text = "\n".join(lines)
    if len(summaries_text) > FINAL_SUMMARIES_MAX_CHARS:
        summaries_text = _compress_summaries(lines, label)

    persons = sorted(merged["persons"].values(), key=lambda e: (-e["mentions"], min(e["pages"] or {0})))
    persons = [{"name": e["name"], "original": sorted(e["original"])[:2], "roles": sorted(e["roles"])[:3],
                "details": e["details"][:3], "pages": sorted(e["pages"])[:15], "mentions": e["mentions"]}
               for e in persons[:MAX_PERSONS]]
    events = sorted(merged["events"], key=lambda e: min(e["pages"] or {0}))[:MAX_EVENTS]
    events = [{**e, "pages": sorted(e["pages"])[:10]} for e in events]
    facts = sorted(merged["facts"], key=lambda f: min(f["pages"] or {0}))[:MAX_FACTS]
    facts = [{**f, "pages": sorted(f["pages"])[:10]} for f in facts]

    material = {
        "document": filename, "unit": label, "units_total": n_units,
        "languages": [k for k, _ in merged["langs"].most_common(5)],
        "summaries": summaries_text, "persons": persons, "events": events, "key_facts": facts,
        "organisations": [k for k, _ in merged["orgs"].most_common(30)],
        "locations": [k for k, _ in merged["locs"].most_common(30)],
        "pictures": merged["images"][:MAX_IMAGES_IN_REPORT],
    }
    prompt = _fill(FINAL_PROMPT, label=label, user_prompt=user_prompt or "None",
                   material=json.dumps(material, ensure_ascii=False, default=list))
    report = _chat_json(prompt, max_tokens=FINAL_MAX_TOKENS)
    if not report.get("overview"):
        report = _fallback_report(merged, persons, events, facts, filename)
    report["_persons_all"] = persons          # full merged list for the frontend / Q&A
    report["_languages"] = material["languages"]
    return report


def _cell(x) -> str:
    return str(x or "").replace("|", "/").replace("\n", " ").strip()


def _render_report(r: dict, label: str, user_prompt: str) -> str:
    out = [f"## {_cell(r.get('title')) or 'Document summary'}"]
    if r.get("document_type"):
        out.append(f"*{_cell(r['document_type'])}*")
    out += ["", "### Overview", str(r.get("overview") or "").strip()]

    kp = [p for p in _as_list(r.get("key_persons")) if isinstance(p, dict)]
    if kp:
        out += ["", "### Key Persons", "| Name | Role | Involvement | Ref |", "|---|---|---|---|"]
        out += [f"| {_cell(p.get('name'))} | {_cell(p.get('role'))} | {_cell(p.get('involvement'))} | "
                f"{_cite(label, _as_list(p.get('pages')))} |" for p in kp]

    ke = [e for e in _as_list(r.get("key_events")) if isinstance(e, dict)]
    if ke:
        out += ["", "### Key Events"]
        for i, e in enumerate(ke, 1):
            extra = "; ".join(x for x in [", ".join(map(str, _as_list(e.get("persons")))),
                                          str(e.get("location") or "")] if x)
            ref = _cite(label, _as_list(e.get("pages")))
            date = f"**{_cell(e.get('date'))}** - " if e.get("date") else ""
            out.append(f"{i}. {date}{_cell(e.get('event'))}"
                       + (f" ({extra})" if extra else "") + (f" *[{ref}]*" if ref else ""))

    det = [x for x in _as_list(r.get("important_details")) if isinstance(x, dict)]
    if det:
        out += ["", "### Important Details"]
        for x in det:
            ref = _cite(label, _as_list(x.get("pages")))
            out.append(f"- {_cell(x.get('detail'))}" + (f" *[{ref}]*" if ref else ""))

    if user_prompt and r.get("answer_to_user_request"):
        out += ["", "### Response to your request", str(r["answer_to_user_request"]).strip()]
    return "\n".join(out)


def _build_corpus(pages: list[dict], results: list[dict]) -> list[dict]:
    """Non-overlapping material used for Q&A: original text + English summaries + picture notes."""
    corpus = []
    for p in pages:
        if p.get("route") == "text":
            t = _page_text(p)
            if t:
                corpus.append({"page": p["page"], "kind": "text", "text": t})
    for r in results:
        data = r.get("data") or {}
        page = r["pages"][0] if r.get("pages") else 0
        for s in _as_list(data.get("page_summaries")):
            if isinstance(s, dict) and s.get("summary"):
                corpus.append({"page": s.get("page") or page, "kind": "summary", "text": str(s["summary"])})
        if r["kind"] == "vision":
            if data.get("key_text_en"):
                corpus.append({"page": page, "kind": "translated_text", "text": str(data["key_text_en"])})
            if data.get("visual_notes"):
                corpus.append({"page": page, "kind": "visual", "text": str(data["visual_notes"])})
        elif r["kind"] == "image" and data.get("description"):
            txt = str(data["description"])
            if data.get("visible_text_en"):
                txt += f" | Visible text: {data['visible_text_en']}"
            corpus.append({"page": page, "kind": "picture", "text": txt})
    order = {"text": 0, "translated_text": 1, "summary": 2, "visual": 3, "picture": 4}

    def key(e):
        try:
            return (int(e["page"]), order.get(e["kind"], 9))
        except (TypeError, ValueError):
            return (0, order.get(e["kind"], 9))
    return sorted(corpus, key=key)


# =====================================================================================
# INDEX (background) - only needed for documents too large to fit in the prompt
# =====================================================================================
def _ensure_index_async(doc_id: str, d: Path) -> None:
    if (_read_json(d / "meta.json", {}) or {}).get("index_ready"):
        _set_status(doc_id, index="ready")
        return
    with _STATUS_LOCK:
        if doc_id in _INDEXING:
            return
        _INDEXING.add(doc_id)
    threading.Thread(target=_build_index, args=(doc_id, d), daemon=True).start()


def _build_index(doc_id: str, d: Path) -> None:
    try:
        _set_status(doc_id, index="building")
        meta = _read_json(d / "meta.json", {}) or {}
        corpus = _read_json(d / "corpus.json", []) or []
        if meta.get("corpus_chars", 0) > FULL_CONTEXT_CHARS:
            chunks = [{"page": e["page"], "kind": e["kind"], "text": piece}
                      for e in corpus for piece in _split(e["text"], EMBED_CHUNK_CHARS, 150)]
            _write_json(d / "chunks.json", chunks)
            vecs = _embed([c["text"] for c in chunks])
            with _qdrant_lock:
                client = _get_qdrant()
                if not client.collection_exists(DOCS_COLLECTION):
                    client.create_collection(DOCS_COLLECTION, vectors_config=VectorParams(
                        size=len(vecs[0]), distance=Distance.COSINE))
                client.delete(DOCS_COLLECTION, points_selector=FilterSelector(filter=_doc_filter(doc_id)))
                pts = [PointStruct(id=str(uuid.uuid5(uuid.NAMESPACE_URL, f"{doc_id}:{i}")), vector=v,
                                   payload={"doc_id": doc_id, **c})
                       for i, (c, v) in enumerate(zip(chunks, vecs))]
                for i in range(0, len(pts), 256):
                    client.upsert(DOCS_COLLECTION, points=pts[i:i + 256])
        meta = _read_json(d / "meta.json", {}) or meta
        meta["index_ready"] = True
        _write_json(d / "meta.json", meta)
        _set_status(doc_id, index="ready")
    except Exception as e:
        _log(d, f"index failed: {traceback.format_exc()}")
        _set_status(doc_id, index="error", index_error=f"{type(e).__name__}: {e}")
    finally:
        with _STATUS_LOCK:
            _INDEXING.discard(doc_id)


_STOP = set(("what which when where who whom whose why how does did is are was were the this that these those "
             "there their about from with into have has had for and but not any all also document page pages "
             "tell show give list please mentioned mention said says").split())


def _keyword_hits(chunks: list[dict], query: str, k: int) -> list[dict]:
    toks = {t.casefold() for t in re.split(r"[\s,.;:!?()\[\]{}\"'।/\\|]+", query) if len(t) >= 3}
    toks = {t for t in toks if t not in _STOP and (not t.isascii() or len(t) >= 4 or any(c.isdigit() for c in t))}
    if not toks:
        return []
    scored = []
    for c in chunks:
        low = c["text"].casefold()
        s = sum(len(t) for t in toks if t in low)
        if s:
            scored.append((s, c))
    scored.sort(key=lambda x: -x[0])
    return [c for _, c in scored[:k]]


def _retrieve(doc_id: str, d: Path, query: str) -> list[dict]:
    hits: list[dict] = []
    try:
        vec = _embed([query], is_query=True)[0]
        with _qdrant_lock:
            res = _get_qdrant().query_points(DOCS_COLLECTION, query=vec, query_filter=_doc_filter(doc_id),
                                             limit=TOP_K, with_payload=True).points
        hits = [{"page": pt.payload.get("page"), "kind": pt.payload.get("kind"),
                 "text": pt.payload.get("text", "")} for pt in res]
    except Exception as e:
        _log(d, f"vector search failed: {e!r}")
    seen = {h["text"] for h in hits}
    for c in _keyword_hits(_read_json(d / "chunks.json", []) or [], query, KEYWORD_K):
        if c["text"] not in seen:
            hits.append(c)
            seen.add(c["text"])

    def key(h):
        try:
            return int(h.get("page") or 0)
        except (TypeError, ValueError):
            return 0
    return sorted(hits, key=key)


def _format_entries(entries: list[dict], label: str, limit: int) -> str:
    out, size = [], 0
    for e in entries:
        block = f"[{_cite(label, [e.get('page')])} | {e.get('kind')}] {e.get('text', '')}"
        if size + len(block) > limit:
            break
        out.append(block)
        size += len(block)
    return "\n\n".join(out)


# =====================================================================================
# SPREADSHEET Q&A - Gemma plans, pandas computes (no eval of generated code)
# =====================================================================================
_OPS = {"==": operator.eq, "!=": operator.ne, ">": operator.gt, "<": operator.lt,
        ">=": operator.ge, "<=": operator.le}
_AGGS = {"count", "sum", "mean", "min", "max", "nunique"}


def _run_table_plan(dfs: dict, plan: dict) -> tuple[str, int]:
    sheet = plan.get("sheet")
    if sheet not in dfs:
        if len(dfs) != 1:
            raise ValueError(f"Unknown sheet {sheet!r}")
        sheet = next(iter(dfs))
    df = dfs[sheet]
    for f in _as_list(plan.get("filters")):
        if not isinstance(f, dict):
            continue
        col, op, val = f.get("column"), f.get("op"), f.get("value")
        if not col and not op:
            continue
        if col not in df.columns:
            raise ValueError(f"Unknown column {col!r}")
        s = df[col]
        if op == "contains":
            mask = s.astype(str).str.contains(str(val), case=False, na=False, regex=False)
        elif op in _OPS:
            if pd.api.types.is_numeric_dtype(s):
                try:
                    val = float(val)
                except (TypeError, ValueError):
                    s, val = s.astype(str), str(val)
            else:
                s, val = s.astype(str), str(val)
            if isinstance(val, str) and op in ("==", "!="):
                s, val = s.str.strip().str.casefold(), val.strip().casefold()
            mask = _OPS[op](s, val)
        else:
            raise ValueError(f"Unknown operator {op!r}")
        df = df[mask]

    gb = [c for c in _as_list(plan.get("group_by")) if c in df.columns]
    agg = plan.get("agg") if isinstance(plan.get("agg"), dict) else {}
    func, acol = agg.get("func"), agg.get("column")
    if func in _AGGS:
        if func == "count":
            res = df.groupby(gb).size().reset_index(name="count") if gb else pd.DataFrame({"count": [len(df)]})
        else:
            if acol not in df.columns:
                raise ValueError(f"Unknown column {acol!r}")
            series = pd.to_numeric(df[acol], errors="coerce") if func in ("sum", "mean") else df[acol]
            if gb:
                res = series.groupby([df[c] for c in gb]).agg(func).reset_index()
            else:
                res = pd.DataFrame({f"{func}_{acol}": [series.agg(func)]})
    else:
        sel = [c for c in _as_list(plan.get("select")) if c in df.columns]
        res = df[sel] if sel else df
    if plan.get("sort_by") in res.columns:
        res = res.sort_values(plan["sort_by"], ascending=bool(plan.get("ascending")))
    try:
        limit = max(1, min(200, int(plan.get("limit") or 20)))
    except (TypeError, ValueError):
        limit = 20
    return res.head(limit).to_csv(index=False), len(res)


def _table_answer_context(d: Path, question: str) -> str:
    try:
        dfs = pd.read_pickle(d / "tables.pkl")
    except Exception:
        return ""
    schema = "\n".join(f"- {name}: {', '.join(df.columns)}" for name, df in dfs.items())
    plan = _chat_json(_fill(TABLE_PLAN_PROMPT, schema=schema, question=question), max_tokens=600)
    if not plan.get("needs_computation"):
        return ""
    try:
        csv, n = _run_table_plan(dfs, plan)
    except Exception as e:
        _log(d, f"table plan failed: {e!r} plan={plan}")
        return ""
    return (f"\nCOMPUTED RESULT (exact, from the full spreadsheet; {n} result rows; "
            f"plan: {json.dumps(plan, ensure_ascii=False)}):\n{csv}\n")


# =====================================================================================
# PUBLIC API
# =====================================================================================
def _run_pipeline(doc_id: str, data: bytes, filename: str, user_prompt: str, progress_cb) -> dict:
    t0 = time.time()
    d = _doc_dir(doc_id)

    def prog(stage, msg, done=0, total=0):
        _set_status(doc_id, state="running", stage=stage, message=msg, done=done, total=total)
        if progress_cb:
            try:
                progress_cb(get_status(doc_id))
            except Exception:
                pass

    cached = _read_json(d / "result.json")
    if cached and cached.get("user_prompt", "") == user_prompt:
        _set_status(doc_id, state="done", stage="done", message="Loaded previous analysis", done=1, total=1)
        _ensure_index_async(doc_id, d)
        return cached

    prog("extracting", "Reading the document")
    pages, label, kind, keeper = _extract(data, filename, d)

    scans = [p for p in pages if p.get("scan")]
    if scans:
        done = 0
        with ThreadPoolExecutor(max_workers=PARALLEL_OCR) as ex:
            futs = {ex.submit(_ocr_page, p): p for p in scans}
            for f in as_completed(futs):
                try:
                    f.result()
                except Exception as e:
                    futs[f]["route"] = "vision"
                    _log(d, f"OCR failed on {label} {futs[f]['page']}: {e!r}")
                done += 1
                prog("ocr", f"Running OCR: {done}/{len(scans)} scanned {label}s", done, len(scans))
    _write_json(d / "pages.json", pages)

    units = _build_units(pages, label)
    results: list[dict] = []
    if units:
        done = 0
        with ThreadPoolExecutor(max_workers=PARALLEL_LLM) as ex:
            futs = [ex.submit(_run_unit, u, label, d) for u in units]
            for f in as_completed(futs):
                try:
                    results.append(f.result())
                except Exception as e:
                    _log(d, f"unit failed: {e!r}")
                done += 1
                prog("analysing", f"Analysing content: {done}/{len(units)}", done, len(units))
    results.sort(key=lambda r: (min(r["pages"]) if r.get("pages") else 0, r["kind"]))

    prog("reporting", "Writing the summary report")
    merged = _merge(results)
    report = _make_report(merged, label, len(pages), filename, user_prompt)
    md = _render_report(report, label, user_prompt)

    corpus = _build_corpus(pages, results)
    _write_json(d / "corpus.json", corpus)
    old_meta = _read_json(d / "meta.json", {}) or {}
    meta = {"doc_id": doc_id, "filename": filename, "kind": kind, "unit_label": label, "units": len(pages),
            "corpus_chars": sum(len(e["text"]) for e in corpus), "created": time.time(),
            "index_ready": old_meta.get("index_ready", False)}
    _write_json(d / "meta.json", meta)
    _write_json(d / "report.json", report)

    stats = {
        "units": len(pages),
        "native": sum(p["source"] == "native" for p in pages),
        "ocr_tesseract": sum(p["source"] == "ocr" and p["route"] == "text" for p in pages),
        "ocr_gemma_vision": sum(p["route"] == "vision" for p in pages),
        "blank_skipped": sum(p["route"] == "skip" for p in pages),
        "pictures_analysed": sum(u["kind"] == "image" for u in units),
        "pictures_over_limit": keeper.skipped_over_limit,
        "failed_parts": len(units) - len(results),
        "seconds": round(time.time() - t0, 1),
    }
    result = {"doc_id": doc_id, "filename": filename, "user_prompt": user_prompt, "unit_label": label,
              "summary_markdown": md, "report": {k: v for k, v in report.items() if not k.startswith("_")},
              "persons": report.get("_persons_all", []), "languages": report.get("_languages", []),
              "stats": stats}
    _write_json(d / "result.json", result)
    _set_status(doc_id, state="done", stage="done", message="Summary ready", done=1, total=1)
    _ensure_index_async(doc_id, d)
    return result


def analyze_document(file_bytes: bytes, filename: str, user_prompt: str | None = None,
                     progress_cb=None, _doc_id: str | None = None) -> dict:
    """Blocking analysis. Returns {doc_id, summary_markdown, report, persons, languages, stats, ...}."""
    doc_id = _doc_id or _make_doc_id(file_bytes)
    try:
        return _run_pipeline(doc_id, file_bytes, filename, (user_prompt or "").strip(), progress_cb)
    except Exception as e:
        _set_status(doc_id, state="error", stage="error", message=f"{type(e).__name__}: {e}")
        _log(_doc_dir(doc_id), traceback.format_exc())
        raise


def start_analysis(file_bytes: bytes, filename: str, user_prompt: str | None = None) -> str:
    """Starts the analysis in a background thread and returns doc_id immediately."""
    doc_id = _make_doc_id(file_bytes)
    with _STATUS_LOCK:
        if doc_id in _RUNNING:
            return doc_id
        _RUNNING.add(doc_id)
    _set_status(doc_id, state="queued", stage="queued", message="Waiting to start", done=0, total=0)

    def job():
        try:
            analyze_document(file_bytes, filename, user_prompt, _doc_id=doc_id)
        except Exception:
            pass  # error is in get_status(doc_id) and log.txt
        finally:
            with _STATUS_LOCK:
                _RUNNING.discard(doc_id)

    threading.Thread(target=job, daemon=True).start()
    return doc_id


def ask_document(doc_id: str, question: str, history: list[dict] | None = None) -> str:
    """Answer a follow-up question. history: [{"role": "user"|"assistant", "content": "..."}]"""
    d = WORK_DIR / _check_id(doc_id)
    meta = _read_json(d / "meta.json")
    if not meta:
        return "This document is no longer available. Please upload it again."
    report = _read_json(d / "report.json", {}) or {}
    label = meta.get("unit_label", "page")

    persons = "; ".join(f"{p.get('name')} ({p.get('role')})" if p.get("role") else str(p.get("name"))
                        for p in _as_list(report.get("key_persons")) if isinstance(p, dict))
    summary = f"{report.get('overview', '')}\nKey persons: {persons}".strip()

    hist_lines = []
    for h in (history or [])[-6:]:
        role = "User" if h.get("role") == "user" else "Assistant"
        hist_lines.append(f"{role}: {str(h.get('content', ''))[:600]}")
    last_user = next((h.get("content", "") for h in reversed(history or []) if h.get("role") == "user"), "")

    computed = _table_answer_context(d, question) if meta.get("kind") == "table" else ""

    corpus = _read_json(d / "corpus.json", []) or []
    if meta.get("corpus_chars", 0) <= FULL_CONTEXT_CHARS:
        context = _format_entries(corpus, label, FULL_CONTEXT_CHARS + 5000)
    elif meta.get("index_ready"):
        query = f"{last_user[:300]} {question}".strip() if last_user else question
        context = _format_entries(_retrieve(doc_id, d, query), label, 40000)
    else:   # index still building -> English summaries only
        light = [e for e in corpus if e["kind"] in ("summary", "translated_text", "picture", "visual")]
        context = ("(Full-text search is still being prepared; answering from summaries.)\n"
                   + _format_entries(light, label, 40000))

    prompt = _fill(QA_PROMPT, filename=meta.get("filename", ""), cite_example=_cite(label, [4]),
                   summary=summary, computed=computed, history="\n".join(hist_lines) or "(none)",
                   question=question, context=context)
    return _chat(prompt, max_tokens=QA_MAX_TOKENS)


def delete_document(doc_id: str) -> None:
    _check_id(doc_id)
    with _qdrant_lock:
        try:
            client = _get_qdrant()
            if client.collection_exists(DOCS_COLLECTION):
                client.delete(DOCS_COLLECTION, points_selector=FilterSelector(filter=_doc_filter(doc_id)))
        except Exception:
            pass
    shutil.rmtree(WORK_DIR / doc_id, ignore_errors=True)
    with _STATUS_LOCK:
        _STATUS.pop(doc_id, None)


def cleanup_old_documents(max_age_days: float = 7) -> int:
    """Deletes analyses older than max_age_days. Returns how many were removed."""
    if not WORK_DIR.exists():
        return 0
    cutoff, n = time.time() - max_age_days * 86400, 0
    for sub in WORK_DIR.iterdir():
        if sub.is_dir() and re.fullmatch(r"[0-9a-f]{16}", sub.name):
            meta = _read_json(sub / "meta.json", {}) or {}
            created = meta.get("created") or sub.stat().st_mtime
            if created < cutoff and sub.name not in _RUNNING:
                delete_document(sub.name)
                n += 1
    return n


if __name__ == "__main__":   # quick manual test: python doc_analyzer.py report.pdf "focus on persons"
    import sys
    if len(sys.argv) < 2:
        print("usage: python doc_analyzer.py <file> [prompt]")
        sys.exit(1)
    f = Path(sys.argv[1])
    res = analyze_document(f.read_bytes(), f.name, sys.argv[2] if len(sys.argv) > 2 else None,
                           progress_cb=lambda s: print(f"[{s['stage']}] {s['message']}"))
    print(res["summary_markdown"])
    print(json.dumps(res["stats"], indent=1))
