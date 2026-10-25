"""
interlinking.py
---------------
Finds hidden links between the KEY PERSONS of a report, up to 3 hops, by searching the
whole cobra database (all categories):

    1 hop  (direct) : A <-> B                      A and B called each other
    2 hops (via 1)  : A <-> X <-> B                both talk to the same person X
    3 hops (via 2)  : A <-> X <-> Y <-> B          A's contact X talks to B's contact Y

Noise control: "hub" numbers (customer care, banks, shops - more than HUB_MAX_CONTACTS
different contacts in cobra), short codes and toll-free numbers are never used as a link.

Relay detection (uses call_date / start_time / end_time / duration): a chain of calls along a
path where each call starts within RELAY_WINDOW_MINUTES after the previous one ended -
e.g. A calls X at 22:00, X calls B at 22:08. That is how a message gets passed on.

Outputs (in <output_parent>/<query>_link_diagrams/):
    interlinks.png   static diagram (goes into the report / network document)
    interlinks.html  interactive diagram - zoom, drag, hover, filters - fully offline
    interlinks.json  every link, path, broker and relay found

Use right after add_network_to_report() in the Generate Report button:
    il = add_interlinks(user_query, docx_path, AppConfig.TARGET_FOLDER, net["network_docx"])
    db.add_message(chat_id, "assistant", f"### Interlinks \\n\\n{il['markdown']}")
    st.session_state.last_interlink_html = il["html"]
"""

from __future__ import annotations

import bisect
import json
import logging
import math
import os
import re
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta
from itertools import combinations

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt                     # noqa: E402
from matplotlib.lines import Line2D                 # noqa: E402

import network_diagram as nd                        # same cobra connection + COLUMNS settings

log = logging.getLogger("interlinking")

# =============================================================================
# CONFIG
# =============================================================================
TIME_COLUMNS = {                    # columns in final_table_voip used for relay detection
    "call_date": "call_date",       # set any of these to None if missing
    "start_time": "start_time",
    "end_time": "end_time",
    "duration": "duration",         # seconds (or HH:MM:SS) - used when end_time is empty
}
HUB_MAX_CONTACTS = 300              # a middle number with more different contacts is a service, not a person
MIN_NUMBER_DIGITS = 8               # shorter numbers (short codes) are never links
IGNORE_PREFIXES = ("1800", "1860")  # toll-free / service prefixes, never links
RELAY_WINDOW_MINUTES = 30
PATHS_PER_PAIR_DRAWN = 3            # in the static image (shortest + strongest first)
PATHS_PER_PAIR_LISTED = 5           # in the tables / interactive diagram
MAX_MIDDLE_NODES_DRAWN = 40         # cap for the static image; the interactive one shows more
SQL_CHUNK = 1000
VIS_JS_PATH = None                  # None = take vis-network.min.js from the installed pyvis package

CLUSTER_COLOURS = ["#C62828", "#1565C0", "#2E7D32", "#6A1B9A", "#EF6C00", "#00838F", "#AD1457", "#4E342E"]
C_MIDDLE, C_BRIDGE, C_RELAY = "#B0BEC5", "#FFB300", "#FF6D00"
C_DIRECT, C_LEG = "#B71C1C", "#546E7A"


# =============================================================================
# DATABASE HELPERS
# =============================================================================
def _col(k):
    return f"`{nd.COLUMNS[k]}`"


def _incat_sql(categories):
    """SQL for 'how many of these calls are in the report categories' + its params."""
    if not categories:
        return "COUNT(*)", []
    ph = ", ".join(["%s"] * len(categories))
    return f"SUM(CASE WHEN {_col('category')} IN ({ph}) THEN 1 ELSE 0 END)", list(categories)


def _chunks(seq, n=SQL_CHUNK):
    seq = list(seq)
    for i in range(0, len(seq), n):
        yield seq[i:i + n]


def _usable(num) -> bool:
    d = re.sub(r"\D", "", str(num or ""))
    return len(d) >= MIN_NUMBER_DIGITS and not d.startswith(IGNORE_PREFIXES)


def _fetch_contacts(conn, number, categories):
    """All calls (all categories) to/from `number`, grouped per (calling, called).
    Exact number first, then last 10 digits."""
    incat, incat_p = _incat_sql(categories)
    base = (f"SELECT {_col('calling_number')}, {_col('called_number')}, COUNT(*), {incat}, {nd._name_cols()} "
            f"FROM `{nd.TABLE}` WHERE ({{cond}}) GROUP BY {_col('calling_number')}, {_col('called_number')}")
    rows = nd._run_query(conn, base.format(cond=f"{_col('calling_number')} = %s OR {_col('called_number')} = %s"),
                         incat_p + [number, number])
    if rows:
        return rows
    last10 = re.sub(r"\D", "", number)[-10:]
    if len(last10) < 10:
        return []
    return nd._run_query(conn, base.format(cond=f"{_col('calling_number')} LIKE %s OR {_col('called_number')} LIKE %s"),
                         incat_p + ["%" + last10, "%" + last10])


def _degrees(conn, raw_numbers):
    """Approximate number of different contacts each number has in the whole of cobra."""
    deg = Counter()
    for chunk in _chunks(raw_numbers):
        ph = ", ".join(["%s"] * len(chunk))
        for a, b in ((_col("calling_number"), _col("called_number")), (_col("called_number"), _col("calling_number"))):
            for num, n in nd._run_query(conn, f"SELECT {a}, COUNT(DISTINCT {b}) FROM `{nd.TABLE}` "
                                              f"WHERE {a} IN ({ph}) GROUP BY {a}", chunk):
                deg[nd._key(num)] += n
    return deg


def _calls_among(conn, raw_numbers, categories):
    """Calls where both sides are in raw_numbers, grouped per pair (with in-category count)."""
    raw_numbers = list(raw_numbers)
    if len(raw_numbers) < 2:
        return []
    incat, incat_p = _incat_sql(categories)
    ph_all = ", ".join(["%s"] * len(raw_numbers))
    rows = []
    for chunk in _chunks(raw_numbers):
        ph = ", ".join(["%s"] * len(chunk))
        rows += nd._run_query(conn, f"SELECT {_col('calling_number')}, {_col('called_number')}, COUNT(*), {incat} "
                                    f"FROM `{nd.TABLE}` WHERE {_col('calling_number')} IN ({ph}) "
                                    f"AND {_col('called_number')} IN ({ph_all}) "
                                    f"GROUP BY {_col('calling_number')}, {_col('called_number')}",
                              incat_p + chunk + raw_numbers)
    return rows


def _call_times(conn, raw_numbers):
    """Individual calls (with times) where both sides are in raw_numbers - for relay detection."""
    tc = {k: v for k, v in TIME_COLUMNS.items() if v}
    if "start_time" not in tc or len(raw_numbers) < 2:
        return []
    cols = ", ".join(f"`{tc[k]}`" if k in tc else "NULL" for k in ("call_date", "start_time", "end_time", "duration"))
    raw_numbers = list(raw_numbers)
    ph_all = ", ".join(["%s"] * len(raw_numbers))
    rows = []
    for chunk in _chunks(raw_numbers):
        ph = ", ".join(["%s"] * len(chunk))
        rows += nd._run_query(conn, f"SELECT {_col('calling_number')}, {_col('called_number')}, {cols} "
                                    f"FROM `{nd.TABLE}` WHERE {_col('calling_number')} IN ({ph}) "
                                    f"AND {_col('called_number')} IN ({ph_all})", chunk + raw_numbers)
    return rows


# =============================================================================
# TIME PARSING (MySQL may return date / datetime / time / timedelta / str)
# =============================================================================
def _to_dt(d, t):
    try:
        if isinstance(t, datetime):
            return t
        if isinstance(d, datetime):
            d = d.date()
        if isinstance(d, str):
            d = datetime.strptime(d.strip()[:10], "%Y-%m-%d").date()
        if isinstance(t, timedelta):
            return datetime.combine(d or date(1970, 1, 1), time()) + t
        if isinstance(t, time):
            return datetime.combine(d or date(1970, 1, 1), t)
        if isinstance(t, str):
            t = t.strip()
            for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
                try:
                    return datetime.strptime(t[:19], fmt)
                except ValueError:
                    pass
            parts = [int(float(x)) for x in t.split(":")]
            parts += [0] * (3 - len(parts))
            return datetime.combine(d or date(1970, 1, 1), time()) + timedelta(hours=parts[0], minutes=parts[1],
                                                                                 seconds=parts[2])
    except Exception:                                          # noqa: BLE001
        return None
    return None


def _duration(v) -> timedelta:
    try:
        if isinstance(v, timedelta):
            return v
        if isinstance(v, str) and ":" in v:
            p = [int(float(x)) for x in v.split(":")]
            p = [0] * (3 - len(p)) + p
            return timedelta(hours=p[0], minutes=p[1], seconds=p[2])
        return timedelta(seconds=float(v or 0))
    except Exception:                                          # noqa: BLE001
        return timedelta(0)


# =============================================================================
# FIND THE LINKS
# =============================================================================
def find_interlinks(key_persons_file: str) -> dict:
    with open(key_persons_file, "r", encoding="utf-8") as fh:
        kp = json.load(fh)
    categories = kp.get("categories", [])
    persons = [p for p in kp.get("key_persons", []) if p.get("numbers")]
    pid_of = {}                                      # number key -> person index
    for i, p in enumerate(persons):
        for n in p["numbers"]:
            pid_of[nd._key(n)] = i
    P = [f"P{i}" for i in range(len(persons))]       # node ids for key persons

    contacts = [dict() for _ in persons]             # per person: contact key -> {calls, in_cat}
    raw_of, names_of = defaultdict(Counter), defaultdict(Counter)
    direct = {}                                      # (i, j) -> {calls, in_cat}

    conn = nd._connect()
    try:
        # ---- 1st level: every contact of every key person (all categories) -----------------------
        for i, p in enumerate(persons):
            own = {nd._key(n) for n in p["numbers"]}
            for number in p["numbers"]:
                for calling, called, cnt, incat, cname, dname in _fetch_contacts(conn, number, categories):
                    kc, kd = nd._key(calling), nd._key(called)
                    if kc in own and kd in own:
                        continue
                    if kc in own:
                        c, c_raw, c_name = kd, called, dname
                    elif kd in own:
                        c, c_raw, c_name = kc, calling, cname
                    else:
                        continue
                    if not c:
                        continue
                    e = contacts[i].setdefault(c, {"calls": 0, "in_cat": 0})
                    e["calls"] += cnt
                    e["in_cat"] += int(incat or 0)
                    raw_of[c][str(c_raw).strip()] += cnt
                    if nd._clean_name(c_name):
                        names_of[c][nd._clean_name(c_name)] += cnt
            for c, e in contacts[i].items():          # direct links to other key persons
                j = pid_of.get(c)
                if j is not None and j != i:
                    pair = tuple(sorted((i, j)))
                    if e["calls"] > direct.get(pair, {"calls": 0})["calls"]:
                        direct[pair] = dict(e)

        # ---- candidate middle numbers: not key persons, not short / toll-free, not hubs ------------
        all_cand = {c for cs in contacts for c in cs if c not in pid_of}
        cand = {c for c in all_cand if _usable(raw_of[c].most_common(1)[0][0] if raw_of[c] else c)}
        service = all_cand - cand                     # short codes / toll-free
        deg = _degrees(conn, sorted({r for c in cand for r in raw_of[c]}))
        hubs = {c for c in cand if deg[c] > HUB_MAX_CONTACTS}
        middle = cand - hubs
        holders = defaultdict(set)                    # middle key -> key persons who talk to it
        for i, cs in enumerate(contacts):
            for c in cs:
                if c in middle:
                    holders[c].add(i)

        # ---- calls between middle numbers (for 3-hop paths) -----------------------------------------
        mm = defaultdict(lambda: {"calls": 0, "in_cat": 0})
        mm_raw = sorted({r for c in middle for r in raw_of[c]})
        for calling, called, cnt, incat in _calls_among(conn, mm_raw, categories):
            a, b = nd._key(calling), nd._key(called)
            if a != b and a in middle and b in middle:
                e = mm[frozenset((a, b))]
                e["calls"] += cnt
                e["in_cat"] += int(incat or 0)

        # ---- build paths -----------------------------------------------------------------------------
        def leg(i, c):
            return {"calls": contacts[i][c]["calls"], "in_cat": contacts[i][c]["in_cat"]}

        paths = defaultdict(list)                     # (i, j) -> [path dict]
        for (i, j), e in direct.items():
            paths[(i, j)].append({"nodes": [P[i], P[j]], "legs": [e], "hops": 1})
        for x, hs in holders.items():                 # A - X - B
            for i, j in combinations(sorted(hs), 2):
                paths[(i, j)].append({"nodes": [P[i], x, P[j]], "legs": [leg(i, x), leg(j, x)], "hops": 2})
        for pair, e in mm.items():                    # A - X - Y - B
            x, y = tuple(pair)
            for a, b in ((x, y), (y, x)):
                for i in holders.get(a, ()):
                    for j in holders.get(b, ()):
                        if i == j or b in contacts[i] or a in contacts[j]:
                            continue                  # same person, or a shorter path already exists
                        lo, hi = (i, j) if i < j else (j, i)
                        nodes = [P[i], a, b, P[j]] if i < j else [P[j], b, a, P[i]]
                        legs = [leg(i, a), dict(e), leg(j, b)] if i < j else [leg(j, b), dict(e), leg(i, a)]
                        paths[(lo, hi)].append({"nodes": nodes, "legs": legs, "hops": 3})
        for pair in paths:
            uniq = {tuple(p["nodes"]): p for p in paths[pair]}
            for p in uniq.values():
                p["strength"] = min(lg["calls"] for lg in p["legs"])
                p["in_report_categories"] = all(lg["in_cat"] > 0 for lg in p["legs"])
            paths[pair] = sorted(uniq.values(), key=lambda p: (p["hops"], -p["strength"]))

        # ---- relays on the listed paths --------------------------------------------------------------
        listed = {pair: ps[:PATHS_PER_PAIR_LISTED] for pair, ps in paths.items()}
        node_raw = set()
        for ps in listed.values():
            for p in ps:
                if p["hops"] >= 2:
                    for n in p["nodes"]:
                        node_raw |= set(persons[int(n[1:])]["numbers"]) if n in P else set(raw_of[n])
        times = defaultdict(list)                     # frozenset(node, node) -> [(start, end)]
        for calling, called, d, st, et, du in _call_times(conn, sorted(node_raw)):
            a, b = nd._key(calling), nd._key(called)
            na = P[pid_of[a]] if a in pid_of else a
            nb = P[pid_of[b]] if b in pid_of else b
            s = _to_dt(d, st)
            if s is None or na == nb:
                continue
            e = _to_dt(d, et) if et not in (None, "") else None
            if e is None or e < s:
                e = s + _duration(du)
            times[frozenset((na, nb))].append((s, e))
    finally:
        conn.close()

    for k in times:
        times[k].sort()
    window = timedelta(minutes=RELAY_WINDOW_MINUTES)

    def relays(nodes):
        legs = [times.get(frozenset((nodes[k], nodes[k + 1])), []) for k in range(len(nodes) - 1)]
        if any(not lg for lg in legs):
            return 0, None
        starts = [[s for s, _ in lg] for lg in legs]
        count, first = 0, None
        for s0, e0 in legs[0]:
            end, ok = e0, True
            for k in range(1, len(legs)):
                idx = bisect.bisect_left(starts[k], end)
                if idx < len(starts[k]) and starts[k][idx] <= end + window:
                    end = legs[k][idx][1]
                else:
                    ok = False
                    break
            if ok:
                count += 1
                first = first or s0
        return count, first

    for ps in listed.values():
        for p in ps:
            if p["hops"] >= 2:
                c, f = relays(p["nodes"])
                p["relays"], p["first_relay"] = c, (f.strftime("%Y-%m-%d %H:%M") if f else None)
            else:
                p["relays"], p["first_relay"] = 0, None

    # ---- brokers (middle numbers that join several key persons) and clusters ----------------------
    broker = defaultdict(lambda: {"persons": set(), "pairs": set(), "strength": 0, "relays": 0})
    for pair, ps in listed.items():
        for p in ps:
            for n in p["nodes"][1:-1]:
                b = broker[n]
                b["persons"] |= set(pair)
                b["pairs"].add(pair)
                b["strength"] += p["strength"]
                b["relays"] += p.get("relays", 0)
    brokers = sorted(({"key": n, "number": raw_of[n].most_common(1)[0][0] if raw_of[n] else n,
                       "name": names_of[n].most_common(1)[0][0] if names_of[n] else "",
                       "links_persons": sorted(b["persons"]), "pairs": len(b["pairs"]),
                       "strength": b["strength"], "relays": b["relays"]} for n, b in broker.items()),
                     key=lambda b: (-len(b["links_persons"]), -b["pairs"], -b["strength"]))

    import networkx as nx
    g = nx.Graph()
    g.add_nodes_from(range(len(persons)))
    g.add_edges_from(paths.keys())
    clusters = [sorted(c) for c in nx.connected_components(g)]
    cluster_of = {i: ci for ci, c in enumerate(sorted(clusters, key=len, reverse=True)) for i in c}
    isolated = [i for i in range(len(persons)) if g.degree(i) == 0]

    return {"query": kp.get("query"), "categories": categories, "persons": persons, "P": P,
            "paths": dict(paths), "listed": listed, "brokers": brokers, "cluster_of": cluster_of,
            "isolated": isolated, "hubs_ignored": len(hubs) + len(service), "middle_checked": len(middle),
            "raw_of": raw_of, "names_of": names_of}


# =============================================================================
# LABELS
# =============================================================================
def _pname(res, i):
    p = res["persons"][i]
    return p.get("name") or p["numbers"][0]


def _node_label(res, n, multiline=True):
    if n in res["P"]:
        i = int(n[1:])
        p = res["persons"][i]
        return (f"{p['name']}\n{p['numbers'][0]}" if multiline else f"{p['name']} ({p['numbers'][0]})") \
            if p.get("name") else p["numbers"][0]
    raw = res["raw_of"][n].most_common(1)[0][0] if res["raw_of"][n] else n
    name = res["names_of"][n].most_common(1)[0][0] if res["names_of"][n] else ""
    return (f"{raw}\n{name}" if multiline else f"{raw} ({name})") if name else raw


def _path_text(res, p):
    return " → ".join(_node_label(res, n, multiline=False) for n in p["nodes"])


def _hop_text(h):
    return {1: "Direct", 2: "Via 1 person", 3: "Via 2 persons"}[h]


# =============================================================================
# GRAPH FOR DRAWING (shared by the static image and the interactive page)
# =============================================================================
def _graph(res, per_pair, max_middle=None):
    chosen = {pair: ps[:per_pair] for pair, ps in res["listed"].items()}
    if max_middle is not None:                         # keep the most important middle nodes only
        imp = Counter()
        for ps in chosen.values():
            for p in ps:
                for n in p["nodes"][1:-1]:
                    imp[n] += 1000 // p["hops"] + p["strength"]
        keep = {n for n, _ in imp.most_common(max_middle)}
        chosen = {pair: [p for p in ps if all(n in keep for n in p["nodes"][1:-1])] for pair, ps in chosen.items()}
    nodes, edges = {}, {}
    for pair, ps in chosen.items():
        for p in ps:
            for n in p["nodes"]:
                nodes.setdefault(n, {"relay": 0, "persons": set()})
            for k, lg in enumerate(p["legs"]):
                a, b = p["nodes"][k], p["nodes"][k + 1]
                e = edges.setdefault(frozenset((a, b)), {"calls": lg["calls"], "in_cat": lg["in_cat"],
                                                         "hops": set(), "direct": p["hops"] == 1})
                e["hops"].add(p["hops"])
            for n in p["nodes"][1:-1]:
                nodes[n]["persons"] |= set(pair)
                nodes[n]["relay"] += p.get("relays", 0)
    for i in range(len(res["persons"])):                # every key person is shown
        nodes.setdefault(res["P"][i], {"relay": 0, "persons": set()})
    hidden = sum(len(ps) for ps in res["listed"].values()) - sum(len(ps) for ps in chosen.values())
    return nodes, edges, hidden


def _layout(res, nodes, edges):
    import networkx as nx
    linked = [i for i in range(len(res["persons"])) if i not in res["isolated"]]
    # ring order: clusters together, and inside a cluster each person next to the one they are most
    # closely linked to (fewest hops, most calls) - keeps lines short and avoids crossings
    closeness = {}
    for (i, j), ps in res["paths"].items():
        closeness[(i, j)] = closeness[(j, i)] = (ps[0]["hops"], -ps[0]["strength"])
    order = []
    for ci in sorted({res["cluster_of"][i] for i in linked}):
        members = [i for i in linked if res["cluster_of"][i] == ci]
        cur = max(members, key=lambda i: sum(1 for j in members if (i, j) in closeness))
        left = set(members) - {cur}
        order.append(cur)
        while left:
            cur = min(left, key=lambda j: closeness.get((cur, j), (9, 0)) + (j,))
            left.discard(cur)
            order.append(cur)
    pos = {}
    for k, i in enumerate(order):
        ang = math.pi / 2 - 2 * math.pi * k / max(len(order), 1)
        pos[res["P"][i]] = (math.cos(ang), math.sin(ang))
    for k, i in enumerate(res["isolated"]):                               # isolated persons: a row below
        pos[res["P"][i]] = (-1.0 + 2.0 * (k + 0.5) / max(len(res["isolated"]), 1), -1.55)
    g = nx.Graph()
    g.add_nodes_from(nodes)
    for e, d in edges.items():
        a, b = tuple(e)
        g.add_edge(a, b, weight=1 + math.log1p(d["calls"]))
    init = dict(pos)
    for n in nodes:
        if n not in init:
            ps = [pos[res["P"][i]] for i in nodes[n]["persons"] if res["P"][i] in pos] or [(0, 0)]
            cx, cy = sum(x for x, _ in ps) / len(ps), sum(y for _, y in ps) / len(ps)
            h = (hash(n) % 1000) / 1000
            init[n] = (cx * 0.6 + 0.15 * math.cos(6.28 * h), cy * 0.6 + 0.15 * math.sin(6.28 * h))
    fixed = [res["P"][i] for i in range(len(res["persons"]))]
    movable = [n for n in nodes if n not in fixed]
    if movable:
        pos = nx.spring_layout(g, pos=init, fixed=fixed, k=0.32, iterations=200, seed=7, weight=None)
    else:
        pos = init
    return pos


# =============================================================================
# STATIC IMAGE (for Word / PDF / chat)
# =============================================================================
def draw_static(res, path):
    nodes, edges, hidden = _graph(res, PATHS_PER_PAIR_DRAWN, MAX_MIDDLE_NODES_DRAWN)
    pos = _layout(res, nodes, edges)
    fig, ax = plt.subplots(figsize=(15, 15))
    ax.axis("off")
    ax.set_xlim(-1.45, 1.45)
    ax.set_ylim(-1.8 if res["isolated"] else -1.45, 1.45)
    max_calls = max((d["calls"] for d in edges.values()), default=1)
    label_edges = len(edges) <= 70

    for e, d in sorted(edges.items(), key=lambda kv: kv[1]["direct"]):
        a, b = tuple(e)
        (x1, y1), (x2, y2) = pos[a], pos[b]
        w = 1.0 + 5.0 * math.log1p(d["calls"]) / math.log1p(max_calls)
        ax.plot([x1, x2], [y1, y2], color=C_DIRECT if d["direct"] else C_LEG, lw=w * (1.4 if d["direct"] else 1),
                ls="-" if d["in_cat"] else (0, (5, 3)), alpha=0.9 if d["direct"] else 0.7, zorder=2 if d["direct"] else 1)
        if label_edges or d["direct"]:
            ax.text((x1 + x2) / 2, (y1 + y2) / 2, str(d["calls"]), fontsize=7.5, ha="center", va="center", zorder=5,
                    bbox=dict(boxstyle="round,pad=0.15", fc="white", ec=C_DIRECT if d["direct"] else C_LEG, lw=0.6))

    for n, d in nodes.items():
        x, y = pos[n]
        if n in res["P"]:
            i = int(n[1:])
            colour = "#9E9E9E" if i in res["isolated"] else CLUSTER_COLOURS[res["cluster_of"][i] % len(CLUSTER_COLOURS)]
            ax.scatter([x], [y], s=2300, c=colour, edgecolors="black", linewidths=1.5, zorder=8)
            ax.text(x, y - 0.075, _node_label(res, n), fontsize=9.5, fontweight="bold", ha="center", va="top", zorder=9,
                    bbox=dict(boxstyle="round,pad=0.25", fc="white", ec=colour, lw=1.2))
        else:
            bridge = len(d["persons"]) >= 3
            ax.scatter([x], [y], s=520 if bridge else 260, c=C_BRIDGE if bridge else C_MIDDLE,
                       marker="D" if bridge else "o", edgecolors=C_RELAY if d["relay"] else "#37474F",
                       linewidths=3.0 if d["relay"] else 0.8, zorder=6)
            ax.text(x, y + 0.04, _node_label(res, n), fontsize=7.5, ha="center", va="bottom", zorder=7)

    n_pairs = len(res["paths"])
    by_hop = Counter(ps[0]["hops"] for ps in res["paths"].values())
    fig.suptitle(f"Interlinks between key persons: {res['query']}", fontsize=16, fontweight="bold", y=0.975)
    sub = (f"{n_pairs} linked pair{'s' if n_pairs != 1 else ''} of key persons  (shortest link: {by_hop[1]} direct, "
           f"{by_hop[2]} via 1 person, {by_hop[3]} via 2 persons)   |   all categories searched   |   "
           f"{res['hubs_ignored']} hub/service numbers ignored")
    if hidden:
        sub += f"   |   {hidden} more paths in the table / interactive diagram"
    fig.text(0.5, 0.948, sub, ha="center", fontsize=9.5, color="#424242")
    if res["isolated"]:
        ax.text(0, -1.4, "No link found (up to 3 hops):", ha="center", fontsize=9, color="#616161", style="italic")

    legend = [Line2D([0], [0], color=C_DIRECT, lw=4, label="Direct calls between key persons"),
              Line2D([0], [0], color=C_LEG, lw=2.5, label="Calls through middle person(s)"),
              Line2D([0], [0], color=C_LEG, lw=2.5, ls=(0, (5, 3)), label="Only in OTHER categories than the report"),
              Line2D([0], [0], marker="o", color="w", markerfacecolor=CLUSTER_COLOURS[0], markersize=13,
                     label="Key person (colour = linked group)"),
              Line2D([0], [0], marker="o", color="w", markerfacecolor=C_MIDDLE, markersize=10, label="Middle person"),
              Line2D([0], [0], marker="D", color="w", markerfacecolor=C_BRIDGE, markersize=11,
                     label="Bridge: links 3+ key persons"),
              Line2D([0], [0], marker="o", color="w", markerfacecolor=C_MIDDLE, markeredgecolor=C_RELAY,
                     markeredgewidth=3, markersize=11, label=f"Relay: calls passed on within {RELAY_WINDOW_MINUTES} min")]
    ax.legend(handles=legend, loc="lower center", bbox_to_anchor=(0.5, -0.07), ncol=3, fontsize=9, frameon=False)
    fig.text(0.5, 0.012, "Numbers on lines = calls between the two. Thicker line = more calls.",
             ha="center", fontsize=8.5, color="#616161")
    fig.savefig(path, dpi=120, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return pos


# =============================================================================
# INTERACTIVE PAGE (offline, vis-network)
# =============================================================================
def _vis_js() -> str | None:
    p = VIS_JS_PATH
    if not p:
        try:
            import importlib.util
            spec = importlib.util.find_spec("pyvis")
            if spec and spec.origin:
                p = os.path.join(os.path.dirname(spec.origin), "templates", "lib", "vis-9.1.2", "vis-network.min.js")
        except Exception:                                        # noqa: BLE001
            p = None
    if p and os.path.isfile(p):
        with open(p, "r", encoding="utf-8") as fh:
            return fh.read()
    return None


_HTML = """<!doctype html><html><head><meta charset="utf-8"><title>Interlinks - __TITLE__</title>
<style>
 body{margin:0;font-family:Segoe UI,Arial,sans-serif;color:#212121;background:#fafafa}
 header{padding:10px 16px;background:#263238;color:#fff} header h1{margin:0;font-size:18px}
 header div{font-size:12px;opacity:.85;margin-top:3px}
 #wrap{display:flex;height:calc(100vh - 62px)} #net{flex:1;background:#fff;border-right:1px solid #ddd}
 #side{width:340px;overflow:auto;padding:12px 14px;font-size:13px} #side h3{margin:14px 0 6px;font-size:14px}
 label{display:block;margin:3px 0;cursor:pointer} #info{white-space:pre-line;background:#f1f3f4;padding:8px;
 border-radius:4px;min-height:60px} .key{display:flex;align-items:center;margin:3px 0}
 .sw{width:22px;height:4px;margin-right:8px} .dot{width:12px;height:12px;border-radius:50%;margin-right:8px}
</style></head><body>
<header><h1>Interlinks between key persons: __TITLE__</h1><div>__SUB__</div></header>
<div id="wrap"><div id="net"></div><div id="side">
 <h3>Show links</h3>
 <label><input type="checkbox" class="hop" value="1" checked> Direct</label>
 <label><input type="checkbox" class="hop" value="2" checked> Via 1 middle person</label>
 <label><input type="checkbox" class="hop" value="3" checked> Via 2 middle persons</label>
 <label><input type="checkbox" id="outside" checked> Links only in other categories (dashed)</label>
 <h3>Details <span style="font-weight:normal;color:#777">(click a person or line)</span></h3><div id="info">-</div>
 <h3>Legend</h3>
 <div class="key"><div class="sw" style="background:#B71C1C;height:6px"></div>Direct calls between key persons</div>
 <div class="key"><div class="sw" style="background:#546E7A"></div>Calls through middle person(s)</div>
 <div class="key"><div class="dot" style="background:#C62828"></div>Key person (colour = linked group)</div>
 <div class="key"><div class="dot" style="background:#B0BEC5"></div>Middle person</div>
 <div class="key"><div class="dot" style="background:#FFB300;border-radius:2px;transform:rotate(45deg)"></div>Bridge (links 3+ key persons)</div>
 <div class="key"><div class="dot" style="background:#B0BEC5;border:3px solid #FF6D00"></div>Relay (calls passed on within __WIN__ min)</div>
 <p style="color:#777">Drag to move people, scroll to zoom. Numbers on lines = calls.</p>
</div></div>
<script>__VISJS__</script>
<script>
const NODES=__NODES__, EDGES=__EDGES__;
for (const it of NODES.concat(EDGES)) { it._text = it.title; const d=document.createElement('div');
  d.style.whiteSpace='pre-line'; d.innerText=it.title; it.title=d; }
const nodes=new vis.DataSet(NODES), edges=new vis.DataSet(EDGES);
function edgeOn(e){ const hops=[...document.querySelectorAll('.hop:checked')].map(x=>+x.value);
  return e.hops.some(h=>hops.includes(h)) && (e.in_cat || document.getElementById('outside').checked); }
const ev=new vis.DataView(edges,{filter:edgeOn});
const nv=new vis.DataView(nodes,{filter:n=>n.key || ev.get().some(e=>e.from===n.id||e.to===n.id)});
const net=new vis.Network(document.getElementById('net'),{nodes:nv,edges:ev},
 {physics:false,interaction:{hover:true,tooltipDelay:120,navigationButtons:true},
  edges:{smooth:{type:'continuous'},font:{size:16,background:'white',strokeWidth:0}},nodes:{font:{size:16}}});
document.querySelectorAll('input').forEach(c=>c.addEventListener('change',()=>{ev.refresh();nv.refresh();}));
net.on('click',p=>{ const id=p.nodes[0]!==undefined?nodes.get(p.nodes[0]):(p.edges[0]!==undefined?edges.get(p.edges[0]):null);
  document.getElementById('info').innerText=id?id._text:'-'; });
</script></body></html>"""


def build_interactive(res, path, static_pos):
    vis = _vis_js()
    if vis is None:
        log.warning("vis-network.min.js not found (install pyvis offline, or set VIS_JS_PATH) - no interactive page")
        return None
    nodes, edges, _ = _graph(res, PATHS_PER_PAIR_LISTED)
    pos = dict(static_pos)
    pos.update({n: xy for n, xy in _layout(res, nodes, edges).items() if n not in pos})
    max_calls = max((d["calls"] for d in edges.values()), default=1)
    vnodes, vedges = [], []
    for n, d in nodes.items():
        x, y = pos[n]
        if n in res["P"]:
            i = int(n[1:])
            colour = "#9E9E9E" if i in res["isolated"] else CLUSTER_COLOURS[res["cluster_of"][i] % len(CLUSTER_COLOURS)]
            p = res["persons"][i]
            pairs = [pr for pr in res["paths"] if i in pr]
            title = (f"KEY PERSON\n{p.get('name') or ''}\n{', '.join(p['numbers'])}\nRole: {p.get('role') or '-'}\n"
                     f"Linked to {len(pairs)} other key person(s)")
            vnodes.append({"id": n, "label": _node_label(res, n), "x": x * 520, "y": -y * 520, "key": True,
                           "shape": "dot", "size": 30, "color": {"background": colour, "border": "#000"},
                           "font": {"size": 22, "face": "Segoe UI, Arial", "bold": True}, "title": title})
        else:
            bridge = len(d["persons"]) >= 3
            who = ", ".join(_pname(res, i) for i in sorted(d["persons"]))
            title = (f"MIDDLE PERSON\n{_node_label(res, n)}\nConnects: {who}"
                     + (f"\nRelays through this person: {d['relay']}" if d["relay"] else ""))
            vnodes.append({"id": n, "label": _node_label(res, n), "x": x * 520, "y": -y * 520, "key": False,
                           "shape": "diamond" if bridge else "dot", "size": 16 if bridge else 10,
                           "borderWidth": 4 if d["relay"] else 1, "font": {"size": 16},
                           "color": {"background": C_BRIDGE if bridge else C_MIDDLE,
                                     "border": C_RELAY if d["relay"] else "#37474F"}, "title": title})
    for k, (e, d) in enumerate(edges.items()):
        a, b = tuple(e)
        w = 1 + 7 * math.log1p(d["calls"]) / math.log1p(max_calls)
        vedges.append({"id": k, "from": a, "to": b, "label": str(d["calls"]), "width": w,
                       "color": {"color": C_DIRECT if d["direct"] else C_LEG, "opacity": 0.85},
                       "dashes": False if d["in_cat"] else [12, 10], "hops": sorted(d["hops"]),
                       "in_cat": bool(d["in_cat"]),
                       "title": f"{_node_label(res, a, False)}\n↔ {_node_label(res, b, False)}\n{d['calls']} calls"
                                + ("" if d["in_cat"] else "\n(none in the report categories)")})
    by_hop = Counter(ps[0]["hops"] for ps in res["paths"].values())
    sub = (f"{len(res['paths'])} linked pairs (shortest: {by_hop[1]} direct, {by_hop[2]} via 1, {by_hop[3]} via 2) "
           f"| all categories | {res['hubs_ignored']} hub/service numbers ignored | report categories: "
           f"{', '.join(res['categories']) or '-'}")
    esc = lambda s: str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    html = (_HTML.replace("__TITLE__", esc(res["query"])).replace("__SUB__", esc(sub))
            .replace("__WIN__", str(RELAY_WINDOW_MINUTES))
            .replace("__NODES__", json.dumps(vnodes, ensure_ascii=False).replace("</", "<\\/"))
            .replace("__EDGES__", json.dumps(vedges, ensure_ascii=False).replace("</", "<\\/"))
            .replace("__VISJS__", vis))
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(html)
    return path


# =============================================================================
# TABLES (markdown - used for chat and Word)
# =============================================================================
def pairs_table(res, per_pair=PATHS_PER_PAIR_LISTED) -> list[str]:
    rows = ["| # | Key person A | Key person B | Link | Path | Calls on each leg | Relays | In report categories |"]
    order = sorted(res["listed"].items(), key=lambda kv: (kv[1][0]["hops"], -kv[1][0]["strength"]))
    k = 0
    for (i, j), ps in order:
        for p in ps[:per_pair]:
            k += 1
            relay = f"{p['relays']} (first {p['first_relay']})" if p.get("relays") else "-"
            rows.append(f"| {k} | {_pname(res, i)} | {_pname(res, j)} | {_hop_text(p['hops'])} | {_path_text(res, p)} | "
                        f"{' / '.join(str(lg['calls']) for lg in p['legs'])} | {relay} | "
                        f"{'Yes' if p['in_report_categories'] else 'No'} |")
        more = len(res["paths"][(i, j)]) - min(per_pair, len(ps))
        if more > 0:
            rows.append(f"|  | {_pname(res, i)} | {_pname(res, j)} | | +{more} more paths in interlinks.json | | | |")
    return rows


def brokers_table(res, top=10) -> list[str]:
    rows = ["| # | Middle person | Connects key persons | Pairs linked | Calls (path strength) | Relays |"]
    for k, b in enumerate(res["brokers"][:top], 1):
        who = ", ".join(_pname(res, i) for i in b["links_persons"])
        rows.append(f"| {k} | {b['number']}{' (' + b['name'] + ')' if b['name'] else ''} | {who} | {b['pairs']} | "
                    f"{b['strength']} | {b['relays'] or '-'} |")
    return rows


def summary_text(res) -> str:
    by_hop = Counter(ps[0]["hops"] for ps in res["paths"].values())
    parts = [f"Links up to 3 hops were searched between {len(res['persons'])} key persons across ALL categories "
             f"in the database. {len(res['paths'])} pair(s) of key persons are linked: {by_hop[1]} directly, "
             f"{by_hop[2]} through one middle person and {by_hop[3]} through two middle persons (shortest link "
             f"counted). {res['hubs_ignored']} hub/service numbers (more than {HUB_MAX_CONTACTS} different "
             f"contacts, short codes, toll-free) were ignored so that they do not create false links."]
    bridges = [b for b in res["brokers"] if len(b["links_persons"]) >= 3]
    if bridges:
        parts.append("Bridges (middle persons linking 3 or more key persons): " +
                     "; ".join(f"{b['number']}{' (' + b['name'] + ')' if b['name'] else ''} links "
                               f"{', '.join(_pname(res, i) for i in b['links_persons'])}" for b in bridges[:5]) + ".")
    rel = sum(p.get("relays", 0) for ps in res["listed"].values() for p in ps)
    if rel:
        parts.append(f"{rel} relay chain(s) found: calls passed along a path within {RELAY_WINDOW_MINUTES} minutes "
                     "of each other, suggesting messages being passed on.")
    if res["isolated"]:
        parts.append("No link found for: " + ", ".join(_pname(res, i) for i in res["isolated"]) + ".")
    return " ".join(parts)


# =============================================================================
# MAIN + WORD / PDF
# =============================================================================
def _json_ready(res, png, html):
    def path_out(p):
        return {"hops": p["hops"], "path": _path_text(res, p), "calls_per_leg": [lg["calls"] for lg in p["legs"]],
                "strength": p["strength"], "in_report_categories": p["in_report_categories"],
                "relays": p.get("relays"), "first_relay": p.get("first_relay")}
    return {"query": res["query"], "categories": res["categories"], "image": png, "interactive": html,
            "hubs_ignored": res["hubs_ignored"], "isolated": [_pname(res, i) for i in res["isolated"]],
            "pairs": [{"a": _pname(res, i), "b": _pname(res, j), "paths": [path_out(p) for p in ps]}
                      for (i, j), ps in res["paths"].items()],
            "brokers": [{**{k: v for k, v in b.items() if k != "links_persons"},
                         "links_persons": [_pname(res, i) for i in b["links_persons"]]} for b in res["brokers"]]}


def generate_interlinks(key_persons_file: str, output_parent: str = nd.LINK_DIAGRAMS_PARENT) -> dict:
    res = find_interlinks(key_persons_file)
    out_dir = os.path.join(output_parent, f"{nd._safe_name(res['query'])}_link_diagrams")
    os.makedirs(out_dir, exist_ok=True)
    png = os.path.join(out_dir, "interlinks.png")
    pos = draw_static(res, png)
    html = build_interactive(res, os.path.join(out_dir, "interlinks.html"), pos)
    with open(os.path.join(out_dir, "interlinks.json"), "w", encoding="utf-8") as fh:
        json.dump(_json_ready(res, png, html), fh, ensure_ascii=False, indent=2)
    res.update({"image": png, "html": html})
    return res


def append_interlinks_section(docx_path: str, res: dict, is_report: bool) -> str | None:
    """New page 'Interlinks': summary, diagram, links table, middle-persons table. Re-makes the PDF."""
    import report_generator as rg
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches
    doc = Document(docx_path)
    doc.add_heading("Interlinks", level=1).paragraph_format.page_break_before = True
    doc.add_paragraph(summary_text(res))
    doc.add_picture(res["image"], width=Inches(5.2))           # fits on the same page as the summary
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if res["paths"]:
        h = doc.add_heading("Links between key persons", level=2)
        h.paragraph_format.page_break_before = True
        rg._flush_table(doc, pairs_table(res))
    if res["brokers"]:
        doc.add_heading("Middle persons linking key persons", level=2)
        rg._flush_table(doc, brokers_table(res))
    doc.save(docx_path)
    total = nd._count_pages(docx_path)
    if is_report:
        nd._set_page_numbers(docx_path, total - 1 if total else None, encl_pages=total - 1 if total else None)
    else:
        nd._set_page_numbers(docx_path, total)
    return rg.export_pdf(docx_path)


def interlinks_markdown(res, max_width=1400) -> str:
    import base64
    import io
    from PIL import Image
    img = Image.open(res["image"]).convert("RGB")
    if img.width > max_width:
        img = img.resize((max_width, int(img.height * max_width / img.width)), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85, optimize=True)
    pic = "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()
    md = [summary_text(res), f"![Interlinks]({pic})"]
    if res["paths"]:
        md += ["**Links between key persons**", "\n".join(pairs_table(res, per_pair=3))]
    if res["brokers"]:
        md += ["**Middle persons linking key persons**", "\n".join(brokers_table(res))]
    if res.get("html"):
        md.append("*An interactive version (zoom, drag, filters) is saved as interlinks.html in the "
                  "link diagrams folder.*")
    return "\n\n".join(md)


def add_interlinks(query: str, report_docx: str, output_parent: str = nd.LINK_DIAGRAMS_PARENT,
                   network_docx: str | None = None) -> dict:
    """
    One call after add_network_to_report(): finds the interlinks, saves interlinks.png / .html / .json,
    adds an 'Interlinks' section to the report (and to the network diagrams document if given) and
    re-makes their PDFs.
    Returns {"image", "html", "report_pdf", "network_pdf", "markdown", "pairs"}.
    """
    import report_generator as rg
    kp_path = rg.key_persons_path(query, os.path.dirname(os.path.abspath(report_docx)))
    out = {"image": None, "html": None, "report_pdf": None, "network_pdf": None, "markdown": "", "pairs": 0}
    if not os.path.isfile(kp_path):
        log.warning("Key persons file not found (%s) - no interlinks", kp_path)
        return out
    res = generate_interlinks(kp_path, output_parent)
    out.update({"image": res["image"], "html": res["html"], "pairs": len(res["paths"]),
                "markdown": interlinks_markdown(res)})
    if len(res["persons"]) >= 2:
        out["report_pdf"] = append_interlinks_section(report_docx, res, is_report=True)
        if network_docx and os.path.isfile(network_docx):
            out["network_pdf"] = append_interlinks_section(network_docx, res, is_report=False)
    return out


# =============================================================================
# STREAMLIT HELPERS for the interactive page
# =============================================================================
def open_interactive(html_path: str) -> None:
    """Opens the interactive diagram in a NEW BROWSER TAB. Works because the Streamlit server and
    the browser run on the same standalone machine."""
    import pathlib
    import webbrowser
    webbrowser.open_new_tab(pathlib.Path(html_path).resolve().as_uri())


def embed_interactive(html_path: str, height: int = 800) -> None:
    """Shows the interactive diagram inside the Streamlit page."""
    import streamlit.components.v1 as components
    with open(html_path, "r", encoding="utf-8") as fh:
        components.html(fh.read(), height=height, scrolling=True)


if __name__ == "__main__":
    import sys
    logging.basicConfig(level=logging.INFO)
    r = generate_interlinks(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else nd.LINK_DIAGRAMS_PARENT)
    print(summary_text(r))
    print("image:", r["image"], "\ninteractive:", r["html"])
