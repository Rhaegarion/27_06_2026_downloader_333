"""
report_generator.py
-------------------
Category-aware narrative intelligence report generator.

Entry point (call this from the Streamlit button handler):

    from report_generator import generate_report
    report_text, docx_path = generate_report("retrieved_records.txt")
    st.markdown(report_text)

Pipeline
    1. parse_records()      .txt -> list of record dicts (no LLM)
    2. build_facts()        deterministic counts / links / pairs (no LLM)
    3. route_report()       picks the template from the category mix
    4. fit_to_context()     condenses very long messages only if needed (LLM)
    5. build_prompt()       shared analytic rules + group-specific structure
    6. call_llm()           Gemma via Ollama
    7. save_docx()          markdown-ish output -> styled .docx

Everything you are likely to tune is in the CONFIG and CATEGORY GROUPS blocks.
"""

from __future__ import annotations

import logging
import math
import os
import re
from collections import Counter, defaultdict
from datetime import datetime

import requests
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

log = logging.getLogger("report_generator")

# =============================================================================
# CONFIG
# =============================================================================
OLLAMA_URL = "http://localhost:11434"
MODEL_NAME = "gemma4:31b"            # exact tag as shown by `ollama list`
OUTPUT_DIR = r"C:\REPORTS_OUTPUT"    # PLACEHOLDER - replace with real path
SYSTEM_NAME = "Automated Report System"
CLASSIFICATION = "RESTRICTED"

TEMPERATURE = 0.3
NUM_PREDICT = 8192          # max tokens the model may write for the report
NUM_CTX_MAX = 65536         # hard ceiling for num_ctx (bounded by your VRAM)
CHARS_PER_TOKEN = 3.0       # conservative estimate (transliterated text is token-heavy)
REQUEST_TIMEOUT = 3600      # seconds
OLLAMA_THINK = None         # set False to explicitly disable Gemma "thinking" if your
                            # Ollama build supports it; None = don't send the flag

CONDENSE_MIN_CHARS = 2500   # only messages longer than this are ever condensed

PRIMARY_THRESHOLD = 0.50    # a group owning >= 50% of records drives the template
SECONDARY_THRESHOLD = 0.25  # a second group with >= 25% gets its focus sections added

# =============================================================================
# CATEGORY GROUPS
# Match is done on a normalised form of the category (lowercase, '&'->'and',
# non-alphanumerics removed). `exact` must match fully, `contains` is substring.
# Order matters: first match wins, so the most severe group is checked first.
# Add your DB's exact spellings to `exact` if substring matching misses any.
# =============================================================================
CATEGORY_GROUPS = {
    "terror": {
        "label": "Insurgency / Terrorism / Hostile Intelligence",
        "exact": {"isi"},
        "contains": ["insurgen", "terror", "militan"],
    },
    "radical": {
        "label": "Fundamentalist / Radical Activity",
        "exact": set(),
        "contains": ["fundamental", "islamic", "radical", "jihad"],
    },
    "military": {
        "label": "Defence / Arakan Army / Arms & Ammunition",
        "exact": set(),
        "contains": ["defence", "defense", "arakan", "arms", "ammunition", "weapon", "military"],
    },
    "crossborder": {
        "label": "Cross-Border Crime (Infiltration / Trafficking / Smuggling / Narcotics)",
        "exact": set(),
        "contains": ["bordercross", "infiltrat", "trafficking", "smuggl", "narcotic", "drug", "cattle"],
    },
    "political": {
        "label": "Diplomatic / Political",
        "exact": set(),
        "contains": ["diplomat", "politic"],
    },
    "situational": {
        "label": "Illegal / Suspicious Activity / Situation Report",
        "exact": set(),
        "contains": ["illegalactivit", "suspicious", "situationreport", "sitrep"],
    },
}
UNMAPPED = "general"

# Groups that always get a dedicated section, even as a small minority.
PRIORITY_GROUPS = ["terror", "radical"]

_JUNK_NAMES = {"", "unknown", "unk", "na", "n/a", "none", "nil", "-", "--", "null", "not known"}

# =============================================================================
# 1. PARSING
# =============================================================================
_FIELD_MAP = {
    "report id": "wan",
    "report date": "date",
    "event id": "event_id",
    "category": "category",
    "calling number": "calling_number",
    "called number": "called_number",
    "calling party": "calling_party",
    "called party": "called_party",
    "subject": "subject",
    "message": "message",
}
_FIELD_RE = re.compile(
    r"^\s*(report id|report date|event id|category|calling number|called number|"
    r"calling party|called party|subject|message)\s*:\s?(.*)$",
    re.IGNORECASE,
)
_RECORD_START_RE = re.compile(r"^\s*report id\s*:", re.IGNORECASE)


def parse_records(source: str) -> list[dict]:
    """`source` is a path to the .txt file, or the raw text itself."""
    if os.path.isfile(source):
        with open(source, "r", encoding="utf-8", errors="replace") as fh:
            text = fh.read()
    else:
        text = source

    blocks, current = [], None
    for line in text.splitlines():
        if _RECORD_START_RE.match(line):
            if current is not None:
                blocks.append(current)
            current = [line]
        elif current is not None:
            current.append(line)
    if current is not None:
        blocks.append(current)

    records, seen = [], set()
    for block in blocks:
        rec = {v: "" for v in _FIELD_MAP.values()}
        in_message, msg_lines = False, []
        for line in block:
            m = _FIELD_RE.match(line)
            if m and not in_message:
                key = _FIELD_MAP[m.group(1).lower()]
                if key == "message":
                    in_message = True
                    msg_lines.append(m.group(2))
                else:
                    rec[key] = m.group(2).strip()
            elif in_message:
                msg_lines.append(line)
        rec["message"] = "\n".join(msg_lines).strip()

        dedupe_key = rec["event_id"] or (rec["wan"], rec["date"], rec["subject"])
        if dedupe_key in seen:
            continue
        seen.add(dedupe_key)
        records.append(rec)

    # chronological order; undated records go last, original order preserved
    records.sort(key=lambda r: (_parse_date(r["date"]) or datetime.max))
    _assign_refs(records)
    return records


def _parse_date(s: str):
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s.strip()[:10], fmt)
        except (ValueError, AttributeError):
            continue
    return None


def _assign_refs(records: list[dict]) -> None:
    """Reference label used in the report. WAN alone unless the WAN repeats in this batch."""
    wan_counts = Counter(r["wan"] for r in records)
    for r in records:
        wan = r["wan"] or "?"
        if wan_counts[r["wan"]] > 1 or not r["wan"]:
            r["ref"] = f"WAN {wan} / Ev {r['event_id'] or '?'}"
        else:
            r["ref"] = f"WAN {wan}"

# =============================================================================
# 2. FACTS (deterministic - the LLM narrates these, it does not recount)
# =============================================================================
_DIAL_CODES = {
    "91": "India", "880": "Bangladesh", "92": "Pakistan", "95": "Myanmar", "977": "Nepal",
    "975": "Bhutan", "86": "China", "94": "Sri Lanka", "93": "Afghanistan", "98": "Iran",
    "971": "UAE", "966": "Saudi Arabia", "974": "Qatar", "968": "Oman", "965": "Kuwait",
    "973": "Bahrain", "60": "Malaysia", "66": "Thailand", "65": "Singapore", "62": "Indonesia",
    "856": "Laos", "90": "Turkey", "44": "UK", "1": "USA/Canada", "49": "Germany", "7": "Russia",
}


def _digits(num: str) -> str:
    return re.sub(r"\D", "", num or "")


def _number_key(num: str) -> str:
    d = _digits(num)
    if d.startswith("00"):
        d = d[2:]
    return d[-10:] if len(d) >= 10 else d   # heuristic: +91 98.. and 98.. collapse together


def _country_of(num: str) -> str | None:
    raw = (num or "").strip()
    d = _digits(raw)
    explicit = raw.startswith("+") or d.startswith("00")
    if d.startswith("00"):
        d = d[2:]
    if not explicit and len(d) <= 10:
        return None                         # looks domestic / no prefix - don't guess
    for n in (3, 2, 1):
        if d[:n] in _DIAL_CODES and 6 <= len(d) - n <= 11:
            return _DIAL_CODES[d[:n]]
    return None


def _clean_name(name: str) -> str:
    n = (name or "").strip()
    return "" if n.lower() in _JUNK_NAMES else n


def _norm_cat(cat: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (cat or "").lower().replace("&", "and"))


def category_group(cat: str) -> str:
    n = _norm_cat(cat)
    for gid, g in CATEGORY_GROUPS.items():
        if n in g["exact"] or any(k in n for k in g["contains"]):
            return gid
    return UNMAPPED


def build_facts(records: list[dict]) -> dict:
    parties = {}                 # key -> {"names": set, "numbers": set, "refs": list}
    pair_refs = defaultdict(list)
    name_numbers = defaultdict(set)
    countries = Counter()

    def touch(num, name, ref):
        name = _clean_name(name)
        key = _number_key(num) or (name.lower() if name else "")
        if not key:
            return None
        p = parties.setdefault(key, {"names": set(), "numbers": set(), "refs": []})
        if name:
            p["names"].add(name)
            if num:
                name_numbers[name.lower()].add(_number_key(num))
        if num:
            p["numbers"].add(num.strip())
        if ref not in p["refs"]:
            p["refs"].append(ref)
        c = _country_of(num)
        if c:
            countries[c] += 1
        return key

    for r in records:
        a = touch(r["calling_number"], r["calling_party"], r["ref"])
        b = touch(r["called_number"], r["called_party"], r["ref"])
        if a and b and a != b:
            pair_refs[tuple(sorted((a, b)))].append(r["ref"])

    dates = [d for d in (_parse_date(r["date"]) for r in records) if d]
    cat_counts = Counter(r["category"] or "Uncategorised" for r in records)
    group_counts = Counter(category_group(r["category"]) for r in records)

    def label(key):
        p = parties[key]
        names = " / ".join(sorted(p["names"])) or "Unidentified"
        nums = ", ".join(sorted(p["numbers"]))
        return f"{names} ({nums})" if nums else names

    ranked = sorted(parties, key=lambda k: -len(parties[k]["refs"]))
    return {
        "total": len(records),
        "start": min(dates).strftime("%Y-%m-%d") if dates else "unknown",
        "end": max(dates).strftime("%Y-%m-%d") if dates else "unknown",
        "span_days": (max(dates) - min(dates)).days + 1 if dates else 0,
        "category_counts": cat_counts,
        "group_counts": group_counts,
        "date_counts": Counter(r["date"] for r in records if r["date"]),
        "top_parties": [(label(k), parties[k]["refs"]) for k in ranked[:12]],
        "linked_parties": [(label(k), parties[k]["refs"]) for k in ranked if len(parties[k]["refs"]) > 1],
        "top_pairs": sorted(((f"{label(a)}  <->  {label(b)}", refs)
                             for (a, b), refs in pair_refs.items()),
                            key=lambda x: -len(x[1]))[:10],
        "multi_number_names": {n: s for n, s in name_numbers.items() if len(s) > 1},
        "countries": countries,
        "unique_parties": len(parties),
    }


def facts_text(f: dict) -> str:
    L = ["COMPUTED FACTS (authoritative - use these figures, do not recount):",
         f"- Total records: {f['total']}",
         f"- Period covered: {f['start']} to {f['end']} ({f['span_days']} days)",
         f"- Distinct parties (by number/name): {f['unique_parties']}",
         "- Records by category: " + ", ".join(f"{c}: {n}" for c, n in f["category_counts"].most_common())]
    busiest = f["date_counts"].most_common(5)
    if busiest and busiest[0][1] > 1:
        L.append("- Busiest dates: " + ", ".join(f"{d} ({n})" for d, n in busiest))
    if f["countries"]:
        L.append("- Party appearances by country (derived from international dialling prefix only; "
                 "numbers without a prefix are not counted): "
                 + ", ".join(f"{c}: {n}" for c, n in f["countries"].most_common()))
    L.append("\nMOST ACTIVE PARTIES (number of records they appear in):")
    L += [f"- {name}: {len(refs)} records [{', '.join(refs)}]" for name, refs in f["top_parties"]]
    if f["top_pairs"]:
        L.append("\nMOST FREQUENT COMMUNICATING PAIRS:")
        L += [f"- {pair}: {len(refs)} records [{', '.join(refs)}]" for pair, refs in f["top_pairs"]]
    if f["linked_parties"]:
        L.append("\nRECORD LINKS (same party appears across these records - use as evidence when linking records):")
        L += [f"- {name}: [{', '.join(refs)}]" for name, refs in f["linked_parties"]]
    if f["multi_number_names"]:
        L.append("\nSAME NAME SEEN ON MULTIPLE NUMBERS (possible number switching, or simply a common name):")
        L += [f"- {n.title()}: {len(s)} numbers" for n, s in f["multi_number_names"].items()]
    return "\n".join(L)


def records_text(records: list[dict]) -> str:
    out = []
    for r in records:
        tag = " [CONDENSED]" if r.get("condensed") else ""
        out.append(
            f"[{r['ref']}] Date: {r['date'] or 'unknown'} | Category: {r['category'] or 'n/a'} | Event ID: {r['event_id'] or 'n/a'}\n"
            f"FROM: {r['calling_party'] or 'Unidentified'} ({r['calling_number'] or 'n/a'})  ->  "
            f"TO: {r['called_party'] or 'Unidentified'} ({r['called_number'] or 'n/a'})\n"
            f"Subject: {r['subject'] or 'n/a'}\n"
            f"Content{tag}:\n{r['message'] or '(empty)'}"
        )
    return "\n\n".join(out)

# =============================================================================
# 3. TEMPLATES
# A section is (TITLE, instruction). Common sections are shared; each group
# lists its full section order and can mix common + specialised sections.
# =============================================================================
BLUF = ("EXECUTIVE SUMMARY",
        "Bottom line up front. One tight paragraph (5-7 sentences) giving the overall picture: what is "
        "happening, who is involved, where, over what period, and why it matters. Then 3-5 bullet-point "
        "KEY JUDGEMENTS, each with a confidence level (High / Moderate / Low) and supporting [WAN] references.")

NETWORK = ("KEY PERSONS & NETWORK",
           "Narrative first: who the central figures are, what role each appears to play, and how records "
           "connect through shared persons or numbers (use RECORD LINKS and COMMUNICATING PAIRS from the "
           "computed facts). Note apparent clusters / separate networks. Link records only where the data "
           "supports it and say what the link is. Then a table:\n"
           "| # | Name | Number(s) | Role / Function | Appears in | Key Contacts | Remarks |\n"
           "Include every person in RECORD LINKS plus any single-record person who is significant.")

CHRONOLOGY = ("CHRONOLOGICAL ANALYSIS",
              "Narrate how events developed over time, grouped into phases or date clusters with a short "
              "sub-heading each (### Phase name - date range). Show escalation, shifts in topic, new actors "
              "appearing, and cause-and-effect between calls. Cite [WAN] refs throughout. End with a compact "
              "timeline table of the KEY developments only (max 15 rows):\n| Date | Ref | Parties | Development |")

OTHER = ("OTHER ACTIVITY",
         "Brief coverage of records outside the main theme of this report (the records whose categories are "
         "listed as secondary below). Summarise them by theme in a short paragraph each and state whether they "
         "connect to the main picture.")

BACKGROUND = ("BACKGROUND CONTEXT (NOT FROM INTERCEPTS)",
              "From your general knowledge only, brief background on organisations, persons of known public "
              "profile, places, or equipment NAMED in the records that a reader would need. Start this section "
              "with the line: '*Source: model general knowledge, not derived from intercepts; may be outdated - "
              "verify before use.*' Keep each entry to 2-4 sentences. If you do not recognise a name, say so - "
              "never guess an identity. Do not repeat this material as intercept findings elsewhere.")


def ASSESSMENT(scale_name: str, scale: str, extra: str = "") -> tuple[str, str]:
    return ("ASSESSMENT",
            f"2-3 paragraphs pulling the picture together. State {scale_name}: {scale}, with justification "
            f"tied to specific [WAN] refs. {extra}Finish with a short 'Intelligence gaps' bullet list: "
            "what is unknown and would change the assessment.")


def ACTIONS(focus: str) -> tuple[str, str]:
    return ("RECOMMENDED ACTIONS",
            "Specific, actionable next steps, grouped as **Immediate**, **Short term**, **Longer term**. "
            f"Each step names the number / person / place it concerns. Emphasis: {focus}")


THREAT = ("THREAT LEVEL", "LOW / MEDIUM / HIGH / CRITICAL")

GROUP_TEMPLATES = {
    # ------------------------------------------------------------------ general
    "general": {
        "title": "Consolidated Intelligence Report",
        "emphasis": "Give a clear, balanced picture across all themes in the records.",
        "sections": [
            BLUF,
            ("SITUATION OVERVIEW",
             "A flowing narrative of the whole body of records: the main storylines, how they relate, and "
             "which matter most. Write so a senior reader understands the picture without reading the records."),
            NETWORK,
            ("THEMATIC ANALYSIS",
             "One ### sub-section per major theme / category group, largest first. For each: what is happening, "
             "who, where, figures involved, and significance. Keep minor themes brief."),
            CHRONOLOGY,
            ASSESSMENT(*THREAT),
            ACTIONS("the highest-risk threads first."),
        ],
    },
    # ------------------------------------------------------------- crossborder
    "crossborder": {
        "title": "Cross-Border Crime Intelligence Report",
        "emphasis": "Geography, routes, crossing points and network identification. Think like a border "
                    "intelligence analyst: where is the border being crossed, how, by whom, and when.",
        "sections": [
            BLUF,
            ("ACTIVITY OVERVIEW",
             "Narrative of what is being moved or who is crossing, the scale, and how the operation works end to end."),
            ("GEOGRAPHIC & ROUTE ANALYSIS",
             "Reconstruct routes from origin -> transit -> crossing point -> destination as described in the calls. "
             "Identify border areas, villages, river/char areas, check-posts, markets or landmarks used for "
             "infiltration or handover. Note timing patterns (night crossings, market days, festivals, weather). "
             "Then a table:\n| Location | Type (Origin / Transit / Crossing point / Destination / Stash) | "
             "Activity | Refs |"),
            ("COMMODITIES, VOLUMES & MONEY",
             "What is being moved (persons, cattle, narcotics, goods), exact quantities, prices, payments and payment "
             "methods as stated. Totals only where figures are explicit. Table:\n"
             "| Item / Persons | Quantity | Price / Payment | Parties | Refs |"),
            NETWORK,
            ("NETWORK LINKAGE ASSESSMENT",
             "Assess whether persons in DIFFERENT records belong to the same network: shared numbers, shared "
             "locations, same middlemen, same code words, same routes, same payment channels. For each suspected "
             "link give the evidence and a confidence level. Identify roles: organiser / financier, coordinator, "
             "linkman / guide, carrier, receiver, and any facilitation by officials or border personnel."),
            ("MODUS OPERANDI & CODE WORDS",
             "Methods, concealment, communication habits, and any coded or euphemistic terms with their likely meaning "
             "(mark meanings as inferred)."),
            CHRONOLOGY,
            ASSESSMENT(*THREAT, extra="Comment on whether activity is increasing, stable or declining. "),
            ACTIONS("crossing points to watch, interdiction windows, numbers to monitor, network nodes to target."),
        ],
    },
    # -------------------------------------------------------------- situational
    "situational": {
        "title": "Situation Report",
        "emphasis": "Clarity above all. The reader must quickly understand WHAT is going on. Write it as a clear "
                    "story, not a list of calls.",
        "sections": [
            ("SITUATION IN BRIEF",
             "One paragraph: what is happening, where, who, since when, and how serious it is. Then 3-5 key "
             "judgements as bullets with confidence levels and [WAN] refs."),
            ("THE SITUATION",
             "The core of this report. Identify the distinct storylines ('threads') in the records. For each thread, "
             "a ### sub-heading and a narrative telling the story from start to latest point: background, what "
             "happened, who did what, and where it stands now. Cite [WAN] refs naturally in the prose."),
            NETWORK,
            ("KNOWN, INFERRED & UNKNOWN",
             "Three short bullet lists: **Established facts** (directly stated in records), **Assessed / inferred** "
             "(your analytic judgement, with confidence), **Unknown** (key open questions)."),
            CHRONOLOGY,
            ("OUTLOOK",
             "Likely developments in the coming days/weeks and the specific indicators that would confirm or "
             "change that outlook."),
            ASSESSMENT(*THREAT),
            ACTIONS("what should be verified or monitored to resolve the unknowns."),
        ],
    },
    # ------------------------------------------------------------------ radical
    "radical": {
        "title": "Radicalisation & Extremist Network Report",
        "emphasis": "Networks, leadership, funding and threat to India. Focus strictly on indicators of "
                    "extremism, radicalisation, unlawful funding or violence; ordinary religious practice, "
                    "worship or community activity is NOT by itself suspicious and must not be presented as such.",
        "sections": [
            BLUF,
            ("ACTIVITY OVERVIEW",
             "Narrative of what is being discussed and done: gatherings, preaching content, recruitment, "
             "propaganda, travel, coordination. Distinguish clearly between extremist indicators and benign activity."),
            ("ORGANISATIONS, LEADERSHIP & ORIGINS",
             "Organisations or groups named or implied, their apparent leaders, who gives direction to whom, and "
             "where leaders/members originate (places, institutions, foreign locations) as stated in the records. "
             "Table:\n| Name | Role / Position | Group | Origin / Base | Refs |"),
            NETWORK,
            ("FUNDING & RESOURCES",
             "Every mention of money or resources: sources (domestic donations, foreign remittance, hawala, trusts/"
             "NGOs, businesses), amounts, channels, intermediaries, purposes. Exact figures only. Table:\n"
             "| Source | Amount | Channel | Recipient / Purpose | Refs |"),
            ("RECRUITMENT, IDEOLOGY & PROPAGANDA",
             "Recruitment methods and targets, ideological themes, literature/media/online channels mentioned."),
            ("FOREIGN LINKAGES",
             "Contacts, funding or direction from outside India, with countries and evidence."),
            ("THREAT ASSESSMENT TO INDIA",
             "Intent, capability and opportunity as evidenced. Any references to targets, dates, events, "
             "mobilisation or violence. State clearly if no violent intent is evidenced."),
            CHRONOLOGY,
            BACKGROUND,
            ASSESSMENT(*THREAT),
            ACTIONS("leaders and funding channels to prioritise, verification of foreign links."),
        ],
    },
    # --------------------------------------------------------------- political
    "political": {
        "title": "Political & Diplomatic Intelligence Report",
        "emphasis": "Actors, positions, intentions and implications for India. Measured, analytical tone; "
                    "separate what was said from what it implies.",
        "sections": [
            BLUF,
            ("KEY ISSUES & DEVELOPMENTS",
             "Narrative of the issues discussed: negotiations, disputes, appointments, elections, visits, agreements, "
             "internal party matters. What changed during the period."),
            ("ACTORS & STAKEHOLDERS",
             "Officials, politicians, parties, ministries, embassies and foreign governments involved; their stated "
             "positions and relationships. Table:\n| Actor | Affiliation / Office | Stated position / Interest | Refs |"),
            NETWORK,
            ("INTENTIONS & POSITIONS",
             "Stated vs apparent intentions; negotiating positions; red lines; divisions within camps; upcoming "
             "meetings, visits, votes or deadlines mentioned."),
            ("IMPLICATIONS FOR INDIA",
             "Security, diplomatic, economic and domestic-political implications, each tied to evidence."),
            CHRONOLOGY,
            ("OUTLOOK & INDICATORS",
             "Most likely course of events and the indicators to watch."),
            BACKGROUND,
            ASSESSMENT("SIGNIFICANCE", "LOW / MEDIUM / HIGH",
                       extra="Also flag any information that appears sensitive, e.g. leaks of official positions. "),
            ACTIONS("information requirements and follow-up collection on key actors."),
        ],
    },
    # ---------------------------------------------------------------- military
    "military": {
        "title": "Military & Arms Intelligence Report",
        "emphasis": "Operational analysis: who, what, when, where, how. Order of battle, commanders, "
                    "activities, movements, weapons and logistics.",
        "sections": [
            BLUF,
            ("OPERATIONAL PICTURE",
             "Narrative of the military / armed-group situation reflected in the records: operations, exercises, "
             "clashes, deployments, arms dealing, and how they relate."),
            ("PERSONNEL & COMMAND",
             "Commanders, officers and key operatives with rank/position, unit/group and location; indicate command "
             "relationships where calls show who instructs whom. Table:\n"
             "| Name | Rank / Position | Unit / Group | Location | Refs |"),
            ("ACTIVITIES & MOVEMENTS",
             "Each exercise, operation, clash or movement: when, where, which units/persons, strength, purpose, outcome. "
             "Table:\n| Date | Activity | Location | Units / Persons | Details | Refs |"),
            ("ARMS, AMMUNITION & LOGISTICS",
             "Weapons, ammunition and equipment named; quantities, prices, suppliers, buyers, payment, delivery routes "
             "and storage. Table:\n| Item | Quantity | Price / Terms | Supplier -> Buyer | Route / Delivery | Refs |\n"
             "For weapon types named, add a brief note on type, origin and capability from general knowledge, "
             "clearly marked *(general knowledge)*."),
            NETWORK,
            ("CAPABILITY & INTENT",
             "What the records indicate about capability (strength, weapons, logistics) and intent (planned "
             "operations, targets). Note implications for India's border regions where relevant."),
            CHRONOLOGY,
            BACKGROUND,
            ASSESSMENT(*THREAT),
            ACTIONS("units, commanders, supply lines and dates to monitor."),
        ],
    },
    # ------------------------------------------------------------------ terror
    "terror": {
        "title": "Terrorism & Insurgency Threat Report",
        "emphasis": "Critical, threat-focused analysis. Identify hierarchy, group formation, handlers and "
                    "operational intent. Surface anything time-sensitive first.",
        "sections": [
            BLUF,
            ("IMMINENT THREAT INDICATORS",
             "FIRST AND MOST IMPORTANT. Any reference to planned attacks, targets, dates, movement of operatives, "
             "weapons/explosives transfer, or instructions to act. Each as a bullet with the exact wording summary "
             "and [WAN] ref. If none are found, write: 'No imminent threat indicators identified in these records.'"),
            ("ORGANISATIONAL STRUCTURE & HIERARCHY",
             "Reconstruct the hierarchy from who gives instructions, who reports, who finances and who executes. "
             "Present as tiers: Leadership / Handlers -> Commanders / Coordinators -> Operatives -> Support "
             "(finance, logistics, shelter, couriers). Show it as an indented list, then a table:\n"
             "| Tier | Name / Alias | Number(s) | Role | Reports to | Confidence | Refs |"),
            NETWORK,
            ("CELL & GROUP FORMATION",
             "Distinct cells or groups, how they are linked, intermediaries/cut-outs, cross-border nodes, and any "
             "signs of new recruitment or restructuring."),
            ("OPERATIONAL ACTIVITY",
             "Planning, training, weapons, finance, recruitment and logistics as evidenced. Exact figures only."),
            ("HANDLER & STATE LINKAGES",
             "Evidence of direction, funding or support from foreign handlers or state agencies (e.g. ISI). "
             "Be precise about what the evidence shows versus what is assessed."),
            ("COMMUNICATIONS SECURITY",
             "Code words, aliases, number switching (see SAME NAME SEEN ON MULTIPLE NUMBERS), use of apps or "
             "intermediaries to avoid detection."),
            CHRONOLOGY,
            BACKGROUND,
            ASSESSMENT(*THREAT, extra="Give a separate confidence level for the threat level. "),
            ACTIONS("imminent threats first, then high-value nodes in the hierarchy, then network disruption."),
        ],
    },
}

# Sections of a group that get pulled into another group's report when that group
# is a strong secondary theme (>= SECONDARY_THRESHOLD).
FOCUS_SECTIONS = {
    "crossborder": ["GEOGRAPHIC & ROUTE ANALYSIS", "NETWORK LINKAGE ASSESSMENT"],
    "situational": ["THE SITUATION"],
    "radical": ["ORGANISATIONS, LEADERSHIP & ORIGINS", "FUNDING & RESOURCES"],
    "political": ["ACTORS & STAKEHOLDERS", "IMPLICATIONS FOR INDIA"],
    "military": ["PERSONNEL & COMMAND", "ARMS, AMMUNITION & LOGISTICS"],
    "terror": ["IMMINENT THREAT INDICATORS", "ORGANISATIONAL STRUCTURE & HIERARCHY"],
}

# =============================================================================
# 4. ROUTING
# =============================================================================
def route_report(facts: dict) -> dict:
    total = max(facts["total"], 1)
    shares = {g: n / total for g, n in facts["group_counts"].items()}
    ranked = sorted(shares.items(), key=lambda x: -x[1])

    primary = ranked[0][0] if ranked and ranked[0][1] >= PRIMARY_THRESHOLD else UNMAPPED
    secondary = [g for g, s in ranked
                 if g not in (primary, UNMAPPED) and s >= SECONDARY_THRESHOLD]
    priority = [g for g in PRIORITY_GROUPS
                if facts["group_counts"].get(g) and g != primary and g not in secondary]

    base = GROUP_TEMPLATES[primary]
    sections = list(base["sections"])
    titles = {t for t, _ in sections}
    insert_at = next(i for i, (t, _) in enumerate(sections) if t == CHRONOLOGY[0])

    extra = []
    for g in secondary:
        for t, instr in GROUP_TEMPLATES[g]["sections"]:
            if t in FOCUS_SECTIONS[g] and t not in titles:
                extra.append((t, instr + f" (Apply this to the {CATEGORY_GROUPS[g]['label']} records.)"))
                titles.add(t)
    for g in priority:
        n = facts["group_counts"][g]
        extra.append((f"PRIORITY FLAG - {CATEGORY_GROUPS[g]['label'].upper()}",
                      f"{n} record(s) in this batch fall under {CATEGORY_GROUPS[g]['label']}. Although a minority, "
                      "they must not be buried: summarise each, state any threat indicators, and say whether they "
                      "connect to persons elsewhere in the report."))

    covered = {primary, *secondary, *priority}
    if primary != UNMAPPED and any(g not in covered for g in facts["group_counts"]):
        extra.append(OTHER)
    sections[insert_at:insert_at] = extra

    title = base["title"]
    emphasis = base["emphasis"]
    if secondary:
        title += " (with " + " & ".join(CATEGORY_GROUPS[g]["label"] for g in secondary) + ")"
        emphasis += " Secondary emphasis: " + " ".join(GROUP_TEMPLATES[g]["emphasis"] for g in secondary)

    return {"primary": primary, "secondary": secondary, "priority": priority,
            "shares": shares, "title": title, "emphasis": emphasis, "sections": sections}

# =============================================================================
# 5. PROMPT
# =============================================================================
SYSTEM_PROMPT = """You are a senior intelligence analyst writing a formal narrative intelligence report from intercepted call records.

ANALYTIC STANDARDS - follow strictly:
1. Every factual claim must come from the records provided. Never invent names, numbers, places, dates, amounts or events.
2. Cite records by their reference label in square brackets exactly as given, e.g. [WAN 123456]. Cite several like [WAN 1; WAN 2].
3. Write in flowing, authoritative narrative prose that tells the story of the records - not a list of call summaries. Tables and bullets only where the structure asks for them.
4. Link records to each other where the data supports it (shared persons, numbers, places, topics, sequence of events) and state the basis for the link. Do not force links that are not there.
5. Numbers must be exact. Use the COMPUTED FACTS for counts, dates and frequencies - do not recount. Quote quantities, amounts and prices exactly as they appear in the records, with units and currency.
6. Clearly separate what the records SAY from your ASSESSMENT. Use calibrated language (confirms / indicates / suggests / possibly) and confidence levels (High / Moderate / Low).
7. Records marked [CONDENSED] are fact extracts of long transcripts; treat them as reliable but less detailed.
8. Background knowledge from outside the records may appear ONLY where a section explicitly allows it, and must be labelled as such.
9. Output GitHub-style markdown: each section heading as '## N. TITLE', sub-headings as '###', tables with '|' pipes and a header separator row. No preamble before the first section and nothing after the last one."""


def build_prompt(records: list[dict], facts: dict, route: dict) -> str:
    groups_present = ", ".join(
        f"{CATEGORY_GROUPS[g]['label'] if g in CATEGORY_GROUPS else 'Other / unmapped'}: "
        f"{facts['group_counts'][g]} records"
        for g, _ in sorted(route["shares"].items(), key=lambda x: -x[1]))

    structure = []
    for i, (title, instr) in enumerate(route["sections"], 1):
        structure.append(f"## {i}. {title}\n{instr}")

    return f"""REPORT TYPE: {route['title']}
ANALYTIC FOCUS: {route['emphasis']}
THEME MIX: {groups_present}

{facts_text(facts)}

==================== RECORDS (chronological) ====================
{records_text(records)}
==================== END OF RECORDS ====================

Write the report now using EXACTLY the following sections, in this order, with these headings.
Scale depth to the evidence: sections with thin evidence stay short; never pad. The chronological
analysis and the group-specific analytic sections should be the most detailed.

{chr(10).join(structure)}"""

# =============================================================================
# 6. LLM
# =============================================================================
def call_llm(system: str, user: str, num_predict: int = NUM_PREDICT) -> str:
    prompt_tokens = (len(system) + len(user)) / CHARS_PER_TOKEN
    num_ctx = min(NUM_CTX_MAX, int(math.ceil((prompt_tokens + num_predict + 1024) / 4096) * 4096))
    payload = {
        "model": MODEL_NAME,
        "messages": [{"role": "system", "content": system},
                     {"role": "user", "content": user}],
        "stream": False,
        "options": {"temperature": TEMPERATURE, "num_ctx": num_ctx, "num_predict": num_predict},
    }
    if OLLAMA_THINK is not None:
        payload["think"] = OLLAMA_THINK
    log.info("LLM call: ~%d prompt tokens, num_ctx=%d", prompt_tokens, num_ctx)
    resp = requests.post(f"{OLLAMA_URL}/api/chat", json=payload, timeout=REQUEST_TIMEOUT)
    resp.raise_for_status()
    return _clean_llm_output(resp.json()["message"]["content"])


def _clean_llm_output(text: str) -> str:
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.DOTALL)
    text = text.strip()
    fence = re.match(r"^```(?:markdown|md)?\s*\n(.*)\n```$", text, flags=re.DOTALL)
    if fence:
        text = fence.group(1).strip()
    # Gemma often writes headings as **1. TITLE** - normalise to '## 1. TITLE'
    text = re.sub(r"^\*\*\s*(\d+\.\s+[^*\n]+?)\s*\*\*\s*$", r"## \1", text, flags=re.MULTILINE)
    text = re.sub(r"^(#{1,4})\s*\*\*(.+?)\*\*\s*$", r"\1 \2", text, flags=re.MULTILINE)
    return text

# =============================================================================
# 4b. CONTEXT FITTING - condense only the longest messages, only if needed
# =============================================================================
CONDENSE_PROMPT = """Extract the intelligence-relevant facts from this intercepted call. Output a compact bullet list (max 200 words).
Keep EXACTLY as stated: names, aliases, phone numbers, places, dates, times, quantities, prices, amounts, weapons, vehicles, code words, instructions and plans.
Do not interpret, assess or add anything not in the text. If the call is routine with nothing of note, say so in one line."""


def fit_to_context(records: list[dict], facts: dict, route: dict) -> list[dict]:
    budget_chars = (NUM_CTX_MAX - NUM_PREDICT - 1024) * CHARS_PER_TOKEN

    def size():
        return len(SYSTEM_PROMPT) + len(build_prompt(records, facts, route))

    if size() <= budget_chars:
        return records
    log.info("Prompt too large (%d chars > %d); condensing longest messages", size(), budget_chars)
    for r in sorted(records, key=lambda r: -len(r["message"])):
        if size() <= budget_chars or len(r["message"]) < CONDENSE_MIN_CHARS:
            break
        header = f"Subject: {r['subject']}\nFROM: {r['calling_party']} TO: {r['called_party']}\n\n"
        r["message"] = call_llm(CONDENSE_PROMPT, header + r["message"], num_predict=600)
        r["condensed"] = True
    if size() > budget_chars:
        log.warning("Still over budget after condensing; hard-truncating remaining long messages")
        per = int(budget_chars * 0.8 / max(len(records), 1))
        for r in records:
            if len(r["message"]) > per:
                r["message"] = r["message"][:per] + " ...[truncated]"
    return records

# =============================================================================
# 7. DOCX
# =============================================================================
def _add_runs(par, text: str):
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        sub = re.split(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", part) if i % 2 == 0 else [part]
        for j, s in enumerate(sub):
            if s:
                run = par.add_run(s)
                run.bold = i % 2 == 1
                run.italic = i % 2 == 0 and j % 2 == 1


def _add_page_number(par):
    run = par.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


def _classification_par(container, text):
    p = container.paragraphs[0] if container.paragraphs else container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    return p


def _flush_table(doc, rows):
    rows = [r for r in rows if not re.fullmatch(r"\|?[\s:\-|]+\|?", r)]
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    if not cells:
        return
    ncols = max(len(c) for c in cells)
    table = doc.add_table(rows=len(cells), cols=ncols)
    table.style = "Table Grid"
    for ri, row in enumerate(cells):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            cell.text = ""
            par = cell.paragraphs[0]
            _add_runs(par, row[ci] if ci < len(row) else "")
            for run in par.runs:
                run.font.size = Pt(9)
                if ri == 0:
                    run.bold = True
    doc.add_paragraph()


def save_docx(report_md: str, path: str) -> str:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    sec = doc.sections[0]
    _classification_par(sec.header, CLASSIFICATION)
    foot = _classification_par(sec.footer, CLASSIFICATION + "    |    Page ")
    _add_page_number(foot)

    table_buf = []
    for line in report_md.splitlines():
        s = line.strip()
        if s.startswith("|"):
            table_buf.append(s)
            continue
        if table_buf:
            _flush_table(doc, table_buf)
            table_buf = []
        if not s or re.fullmatch(r"[-*_]{3,}", s):
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).replace("**", ""), level=0 if level == 1 else level - 1)
        elif re.match(r"^[-*•]\s+", s):
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*•]\s+", "", s))
        elif re.match(r"^\s{2,}[-*•]\s+", line):
            _add_runs(doc.add_paragraph(style="List Bullet 2"), re.sub(r"^\s*[-*•]\s+", "", line))
        else:
            _add_runs(doc.add_paragraph(), s)
    if table_buf:
        _flush_table(doc, table_buf)

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    doc.save(path)
    return path

# =============================================================================
# ENTRY POINT
# =============================================================================
def report_header(facts: dict, route: dict) -> str:
    cats = ", ".join(f"{c} ({n})" for c, n in facts["category_counts"].most_common())
    return (f"**{CLASSIFICATION}**\n\n"
            f"# {route['title'].upper()}\n\n"
            f"**Date:** {datetime.now().strftime('%d %B %Y')}  \n"
            f"**Records analysed:** {facts['total']}  \n"
            f"**Period:** {facts['start']} to {facts['end']}  \n"
            f"**Categories:** {cats}  \n"
            f"**System:** {SYSTEM_NAME}\n\n---\n")


def generate_report(source: str, output_dir: str = OUTPUT_DIR) -> tuple[str, str]:
    """
    source      path to the retrieved-records .txt file (or its raw text)
    output_dir  folder for the .docx
    returns     (report_markdown, docx_path)
    """
    records = parse_records(source)
    if not records:
        raise ValueError("No records could be parsed from the input.")
    facts = build_facts(records)
    route = route_report(facts)
    log.info("Routing: primary=%s secondary=%s priority=%s",
             route["primary"], route["secondary"], route["priority"])

    records = fit_to_context(records, facts, route)
    body = call_llm(SYSTEM_PROMPT, build_prompt(records, facts, route))
    report = report_header(facts, route) + "\n" + body

    fname = f"Report_{route['primary']}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    docx_path = save_docx(report, os.path.join(output_dir, fname))
    return report, docx_path


if __name__ == "__main__":
    import sys
    logging.basicConfig(level=logging.INFO)
    text, path = generate_report(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else OUTPUT_DIR)
    print(text)
    print(f"\nSaved: {path}")
