"""
summary_generator.py
--------------------
Short, dense, fact-only summary of the same inputs used by report_generator.py.
Reuses its parsing, WAN filtering, routing, LLM call and Word helpers.

    from summary_generator import generate_summary
    summary_text, docx_path, pdf_path = generate_summary(context_path, query, json_path)

Saves <query>_summary.docx and <query>_summary.pdf in OUTPUT_DIR (from report_generator).
"""

from __future__ import annotations

import logging
import os
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from report_generator import (
    CATEGORY_GROUPS, DOC_FONT, FOOTER_FONT_SIZE, FOOTER_LEFT, OUTPUT_DIR,
    _add_field, _header_classification, _render_markdown, _right_tab, _setup_page,
    build_facts, call_llm, export_pdf, facts_text, fit_to_context, load_wan_list,
    output_paths, parse_records, records_text, route_report,
)

log = logging.getLogger("summary_generator")

SUMMARY_NUM_PREDICT = 2500     # safety cap; the prompt itself asks for ~200-350 words + table
SUMMARY_FONT_SIZE = 10
SUMMARY_TITLE_SIZE = 16

# What each category group's summary prioritises, and the persons-table columns.
SUMMARY_FOCUS = {
    "general": ("What is happening, who is involved, where, and the key figures.",
                "| Name | Number(s) | Role | Location / Links | Key details |"),
    "crossborder": ("Routes, crossing points and border areas; what is moved (persons, cattle, narcotics, "
                    "goods) with exact quantities, prices and payment; who does what (organiser, linkman, "
                    "carrier, receiver); timings of crossings.",
                    "| Name | Number(s) | Role | Location / Area operated | Key details (goods, amounts, contacts) |"),
    "situational": ("What happened, in what sequence, where, who was involved, and how things stand at the "
                    "latest record.",
                    "| Name | Number(s) | Role | Location | Key details |"),
    "radical": ("Organisations and leaders, origins, funding (amounts, sources, channels), recruitment, "
                "foreign links and any threat indicators. Ordinary religious activity is not by itself "
                "suspicious and must not be presented as such.",
                "| Name | Number(s) | Role / Position | Group | Origin / Base | Key details (funds, links) |"),
    "political": ("Actors, their offices and stated positions, issues discussed, upcoming meetings, visits "
                  "or deadlines, and implications for India.",
                  "| Name | Number(s) | Office / Affiliation | Stated position / Interest | Key details |"),
    "military": ("Commanders and units, locations, movements and exercises (when, where, who), weapons and "
                 "ammunition with quantities, prices, suppliers and buyers.",
                 "| Name | Number(s) | Rank / Position | Unit / Group | Location | Key details (weapons, orders) |"),
    "terror": ("Any threat indicators FIRST (targets, dates, movement of operatives, weapons/explosives); then "
               "hierarchy (who instructs whom), handlers and foreign links, operatives, finance and weapons.",
               "| Name / Alias | Number(s) | Role / Tier | Reports to | Location | Key details |"),
}

SUMMARY_SYSTEM_PROMPT = """You write short, dense intelligence summaries of intercepted call records for a senior officer who wants the picture quickly. It must be short enough to read in a couple of minutes, yet carry the details that matter.

RULES - follow strictly:
1. Facts only, taken from the records. Never invent names, numbers, places, dates, amounts or events. No assessment, risk rating, background knowledge, recommendations or commentary.
2. RELEVANCE: include what answers the USER QUERY and what the PRIORITISE line asks for, with exact names, numbers, places, dates, quantities and amounts. Leave out routine or irrelevant chatter. Say each fact once - merge repeated information.
3. PEOPLE ARE NEVER DROPPED: every individual involved must appear, with their details (all numbers, aliases, role, location, contacts, what they handle). Persons with nothing significant beyond being in contact go together in ONE line after the table: 'Also in contact: Name (number), Name (number) ...'.
4. Cite record references (exact label, e.g. [WAN CEC-KOL 0382/09]) as RARELY as possible: only where it pins down a key fact, never the same reference twice, grouped where one statement covers several.
5. Numbers exact, with units and currency, as written in the records. Use the COMPUTED FACTS for counts.
6. LENGTH: the prose (everything except the table) should be about 200-350 words; table cells terse (a few words each). Shorter is better if nothing is lost.
7. Layout - NO headings or sub-headings of any kind except the single TITLE line:
   a) Narrative, one or two short paragraphs. Its FIRST sentence directly answers the user query in one line (who / what / where / when); the rest tells how things developed, in chronological flow, and where they stand at the latest record.
   b) A one-line lead-in sentence, then the persons table using the columns given. One row per significant person; places tied to a person go in their row.
   c) The 'Also in contact:' line, if needed.
   d) Places or routes not tied to anyone: ONE line, routes written as A -> B -> C. Skip if none.
8. The VERY FIRST line of your output must be: TITLE: Summary on <topic> - <topic> is the user query rephrased minimally into a clean title phrase.
9. Output GitHub-style markdown. Tables use '|' pipes with a header separator row. Nothing after the last line."""


def build_summary_prompt(records: list[dict], facts: dict, route: dict, query: str = "") -> str:
    focus, columns = SUMMARY_FOCUS.get(route["primary"], SUMMARY_FOCUS["general"])
    for g in route["secondary"]:
        focus += " Also: " + SUMMARY_FOCUS[g][0]
    flags = ""
    if route["priority"]:
        flags = ("\nMINORITY RECORDS THAT MUST STILL BE COVERED: "
                 + ", ".join(CATEGORY_GROUPS[g]["label"] for g in route["priority"])
                 + " - include their facts and any threat indicator, even if briefly.")
    return f"""USER QUERY (the request these records were retrieved for): "{query or 'not provided'}"
PRIORITISE IN THIS SUMMARY: {focus}{flags}
PERSONS TABLE COLUMNS: {columns}

{facts_text(facts)}

==================== RECORDS (chronological) ====================
{records_text(records)}
==================== END OF RECORDS ====================

Write the summary now, following the rules and layout exactly. Short and dense; every individual included."""


def _load_filtered_records(source: str, wan_list) -> list[dict]:
    wan_filter = load_wan_list(wan_list) if wan_list is not None else None
    records = parse_records(source, wan_filter)
    if wan_filter is not None and not records and any("|" in k for k in wan_filter):
        log.warning("No record matched on WAN No. + date; retrying on WAN No. only")
        wan_filter = {k.split("|")[0] for k in wan_filter}
        records = parse_records(source, wan_filter)
    if not records:
        raise ValueError("No records left to summarise: none parsed, or none matched the WAN No. list.")
    return records


def save_summary_docx(summary_md: str, path: str) -> str:
    """Single section: SECRET/M.IMMDT header, footer 'Generated by ...' + 'Page X of Y'. No cover."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = DOC_FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), DOC_FONT)
    style.font.size = Pt(SUMMARY_FONT_SIZE)
    style.paragraph_format.space_after = Pt(4)
    doc.styles["Title"].font.size = Pt(SUMMARY_TITLE_SIZE)
    sec = doc.sections[0]
    _setup_page(sec)
    sec.left_margin = sec.right_margin = Inches(0.75)               # compact page
    sec.top_margin = sec.bottom_margin = Inches(0.75)
    _header_classification(sec)

    p = sec.footer.paragraphs[0]
    _right_tab(p, sec)
    r = p.add_run(FOOTER_LEFT + "\t")
    r.font.size = Pt(FOOTER_FONT_SIZE)
    for txt, fld in (("Page ", "PAGE"), (" of ", "NUMPAGES")):
        r = p.add_run(txt)
        r.font.size = Pt(FOOTER_FONT_SIZE)
        _add_field(p, fld, size=FOOTER_FONT_SIZE)

    _render_markdown(doc, summary_md)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    doc.save(path)
    return path


def generate_summary(source: str, query: str, wan_list=None,
                     output_dir: str = OUTPUT_DIR) -> tuple[str, str, str | None]:
    """
    Same inputs as generate_report().
    returns (summary_markdown, docx_path, pdf_path)   pdf_path is None if no PDF converter
    """
    records = _load_filtered_records(source, wan_list)
    facts = build_facts(records)
    route = route_report(facts)
    records = fit_to_context(records, facts, route, query)
    body = call_llm(SUMMARY_SYSTEM_PROMPT, build_summary_prompt(records, facts, route, query),
                    num_predict=SUMMARY_NUM_PREDICT)

    m = re.match(r"^\s*\**\s*TITLE\s*:\s*\**\s*(.+?)\s*\**\s*$", body.split("\n", 1)[0], re.IGNORECASE)
    if m:
        title = m.group(1).strip().strip('"')
        body = body.split("\n", 1)[1].lstrip() if "\n" in body else ""
    else:
        q = re.sub(r"[?.!\s]+$", "", (query or "").strip())
        title = f"Summary on {q[:1].upper() + q[1:]}" if q else "Summary"
    body = re.sub(r"^#{1,6}\s+.*\n?", "", body, flags=re.MULTILINE)   # enforce: no sub-headings
    body = re.sub(r"(^[^|\n].*)\n(?=\|)", r"\1\n\n", body, flags=re.MULTILINE)  # blank line before tables
    body = re.sub(r"(^\|.*)\n(?=[^|\n])", r"\1\n\n", body, flags=re.MULTILINE)   # and after tables
    meta = f"*{facts['total']} intercepts, {facts['start']} to {facts['end']}*"
    summary = f"# {title}\n\n{meta}\n\n{body.strip()}\n"

    docx_path, _ = output_paths(query, "summary", output_dir)
    save_summary_docx(summary, docx_path)
    pdf_path = export_pdf(docx_path)
    return summary, docx_path, pdf_path


if __name__ == "__main__":
    import sys
    logging.basicConfig(level=logging.INFO)
    # usage: python summary_generator.py records.txt "user query" refs_<session_id>.json [output_dir]
    text, d, p = generate_summary(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None,
                                  sys.argv[4] if len(sys.argv) > 4 else OUTPUT_DIR)
    print(text)
    print(f"\nSaved: {d}\n       {p}")
